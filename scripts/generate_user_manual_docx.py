"""Generate DAMS User Manual Word document with role-based screen guides."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "docs" / "DAMS_User_Manual.docx"
SHOT_DIR = ROOT / "docs" / "user-manual-screenshots"


def set_run_font(run, size: int = 12, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 14 if level == 2 else 12, bold=True)


def add_paragraph(doc: Document, text: str, size: int = 12) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, size=size)


def add_bullets(doc: Document, items: list[str], size: int = 12) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_role_badge(doc: Document, role: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"Access level: {role}")
    set_run_font(run, size=11, bold=True)


def add_figure(doc: Document, figure_no: int, caption: str, image_name: str) -> None:
    image_path = SHOT_DIR / image_name
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(4)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(f"Figure {figure_no}: {caption}")
    set_run_font(cap_run, size=11, bold=True)

    if image_path.exists():
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.space_after = Pt(8)
        pic.add_run().add_picture(str(image_path), width=Inches(6.2))
    else:
        add_paragraph(doc, f"[Screenshot not found: {image_name}]", size=10)


def add_screen_section(
    doc: Document,
    figure_no: int,
    title: str,
    role: str,
    image_name: str,
    overview: str,
    steps: list[str] | None = None,
    notes: str | None = None,
) -> None:
    add_heading(doc, title, level=3)
    add_role_badge(doc, role)
    add_paragraph(doc, overview)
    if steps:
        add_bullets(doc, steps)
    if notes:
        add_paragraph(doc, notes)
    add_figure(doc, figure_no, title, image_name)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("User Manual")
    set_run_font(title_run, size=18, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("Deen Associate Management System (DAMS)")
    set_run_font(sub_run, size=14, bold=True)

    doc.add_paragraph()

    add_heading(doc, "1. Introduction", level=1)
    add_paragraph(
        doc,
        "The Deen Associate Management System (DAMS) is a web-based real estate management "
        "platform used to publish projects, manage property units, handle customer booking "
        "requests, confirm sales, track installment schedules, record payments, manage "
        "employees, and monitor company finances. The same application serves two primary "
        "audiences: clients (customers) who browse projects and track their own bookings, "
        "and administrators who operate the full back-office workflow.",
    )
    add_paragraph(
        doc,
        "This manual explains how to use DAMS screen by screen. Each section clearly states "
        "whether the screen is for clients, administrators, or both. Screens marked "
        "Shared are visible to visitors and logged-in users; admin-only screens require "
        "an account with the Admin role.",
    )

    add_heading(doc, "2. User Roles", level=1)
    add_heading(doc, "2.1 Client (Customer)", level=2)
    add_paragraph(
        doc,
        "A client account is created through self-registration. After login, the navigation "
        "bar shows Home, Projects, About, Contact, and My Projects. Clients can browse "
        "available units, submit booking requests, and—once a booking is confirmed—view "
        "their payment progress, installment schedule, application form, and receipts from "
        "My Projects.",
    )
    add_bullets(
        doc,
        [
            "Browse projects and units without administrative controls.",
            "Submit booking requests for available units.",
            "Track confirmed bookings, payments, and documents in My Projects.",
            "Cannot access Requests, Bookings (admin list), Customers, Employees, or Finance.",
        ],
    )

    add_heading(doc, "2.2 Administrator", level=2)
    add_paragraph(
        doc,
        "An administrator account has full operational access. After login, the navigation "
        "bar expands to include Requests, Bookings, Customers, Employees, and Finance in "
        "addition to the public pages. Admins review booking requests, confirm bookings, "
        "generate installment plans, record payments, print official forms and receipts, "
        "manage staff attendance and salaries, and maintain the finance ledger.",
    )
    add_bullets(
        doc,
        [
            "Review and approve or reject booking requests.",
            "Create and manage confirmed bookings and installment schedules.",
            "Record booking and installment payments with printable receipts.",
            "Manage customers, employees, attendance, salaries, and finance entries.",
            "Upload and maintain project media and unit inventory.",
        ],
    )

    add_heading(doc, "3. Getting Started", level=1)
    add_heading(doc, "3.1 Opening the Application", level=2)
    add_paragraph(
        doc,
        "Open the application URL provided by your organization (for example, "
        "http://localhost:5173 during development). The home page introduces Deen Associate "
        "and provides quick links to Projects and Contact.",
    )
    add_screen_section(
        doc,
        1,
        "Home Page",
        "Guest / Client / Admin (public landing page)",
        "01-home.png",
        "The home page welcomes users to Deen Associate and highlights the main entry "
        "points. Use View Projects to browse inventory, or Contact Us to reach the sales "
        "team. The top navigation remains available on every page.",
        steps=[
            "Open the application in a modern browser (Chrome, Edge, or Firefox recommended).",
            "Click View Projects to browse available real estate projects.",
            "Use Log in or Sign Up in the top-right corner when you need an account.",
        ],
    )

    add_heading(doc, "3.2 Registration and Login", level=2)
    add_paragraph(
        doc,
        "New customers create a client account using Sign Up. Existing users sign in with "
        "Log in. Administrator accounts are typically created by system staff and are not "
        "available through public self-registration.",
    )
    add_screen_section(
        doc,
        2,
        "Login Dialog",
        "Guest → Client or Admin after successful authentication",
        "05-login-modal.png",
        "The login dialog collects email and password. On success, the system stores a "
        "secure session token and loads the navigation appropriate to the user's role.",
        steps=[
            "Click Log in on the navigation bar.",
            "Enter your registered email address and password.",
            "Submit the form to access role-based pages.",
        ],
    )
    add_screen_section(
        doc,
        3,
        "Sign Up Dialog",
        "Guest → new Client account",
        "06-signup-modal.png",
        "The sign-up dialog creates a new customer account. Provide full name, email, and "
        "password. After registration, return to Log in with the same credentials.",
        steps=[
            "Click Sign Up on the navigation bar.",
            "Enter full name, email, and password.",
            "Complete registration, then log in with the new account.",
        ],
    )

    add_heading(doc, "3.3 Theme and Profile Controls", level=2)
    add_paragraph(
        doc,
        "The sun/moon icon in the navigation bar toggles light and dark mode. When logged "
        "in, the profile chip shows the account initial and email prefix; the adjacent "
        "button logs the user out.",
    )

    add_heading(doc, "4. Client User Guide", level=1)
    add_paragraph(
        doc,
        "The following screens support the customer journey: discovering projects, "
        "reviewing units, submitting booking interest, and tracking confirmed purchases.",
    )

    add_screen_section(
        doc,
        4,
        "Projects Listing",
        "Guest / Client / Admin (browse mode)",
        "02-projects.png",
        "The Projects page lists all published developments. Each card shows the project "
        "name, location, status, and summary information. Select a project to open its "
        "detail view.",
        steps=[
            "Click Projects in the navigation bar.",
            "Review available developments and their current status.",
            "Open a project to inspect units and media.",
        ],
    )

    add_screen_section(
        doc,
        5,
        "Project Detail — Overview Tab",
        "Guest / Client / Admin",
        "07-project-overview.png",
        "The Overview tab describes the selected project, including location, launch date, "
        "and current delivery stage on the timeline (Planning, Construction, Handover, "
        "Completed). This helps clients understand project progress before booking.",
        steps=[
            "Open a project from the Projects list.",
            "Read the About This Project section and metadata row.",
            "Review the Project Timeline to see the current stage.",
        ],
    )

    add_screen_section(
        doc,
        6,
        "Project Detail — Units Tab",
        "Guest / Client / Admin",
        "08-project-units.png",
        "The Units tab displays all units in the project as cards with photos, type, floor, "
        "area, and availability status (for example Available, Reserved, OnPaymentPlan). "
        "Clients use this screen to choose a unit before submitting a booking request.",
        steps=[
            "Switch to the Units tab on a project page.",
            "Use search and filters for type, status, floor, or sort order.",
            "Open a unit card to view full details and start a booking request.",
        ],
    )

    add_screen_section(
        doc,
        7,
        "Project Detail — Media Tab",
        "Guest / Client / Admin",
        "09-project-media.png",
        "The Media tab shows gallery images and uploaded files for the project. Clients "
        "use this section to review marketing visuals, floor plans, or brochures before "
        "making a purchase decision.",
        steps=[
            "Switch to the Media tab.",
            "Browse uploaded gallery items and file details.",
            "Use carousel controls where available for full-size viewing.",
        ],
    )

    add_paragraph(
        doc,
        "From a unit detail page, an authenticated client can submit a booking request. "
        "The request enters the admin Requests queue for review. Once approved and "
        "converted into a confirmed booking, the unit appears under My Projects.",
    )

    add_screen_section(
        doc,
        8,
        "My Projects",
        "Client only",
        "10-my-projects.png",
        "My Projects is the client dashboard for confirmed bookings. Each row or card "
        "shows booking reference, project, unit, agreed price, booking amount progress, "
        "installment totals, and current status (for example PaymentPlanActive).",
        steps=[
            "Log in with a client account.",
            "Click My Projects in the navigation bar.",
            "Select a booking to view payment history, schedule, forms, and receipts.",
        ],
        notes=(
            "Administrators are redirected away from My Projects because they manage "
            "bookings from the admin Bookings module instead."
        ),
    )

    add_screen_section(
        doc,
        9,
        "Application Form (Client View)",
        "Client (own booking) · Admin (any booking)",
        "14-application-form.png",
        "The application form is the official customer record for a confirmed booking. "
        "It includes serial number, date, unit details, applicant information, and next "
        "of kin fields. Clients access it from their booking detail; administrators open "
        "it via Print Application Form on the admin booking page.",
        steps=[
            "Open the booking from My Projects (client) or Confirmed Bookings (admin).",
            "Choose Print Application Form or the equivalent view action.",
            "Use Print / Save as PDF in the browser for a downloadable copy.",
        ],
    )

    add_screen_section(
        doc,
        10,
        "Payment Receipt (Client View)",
        "Client (own payments) · Admin (all payments)",
        "15-payment-receipt.png",
        "Every recorded payment generates a printable receipt with receipt number, amount "
        "in words, payment type (Down Payment or Installment), mode of payment, and "
        "reference details. Clients can open receipts linked to their bookings; "
        "administrators open them after recording a payment.",
        steps=[
            "Open the booking or installment history.",
            "Select the receipt for a completed payment.",
            "Print or save the receipt as PDF for your records.",
        ],
    )

    add_screen_section(
        doc,
        11,
        "About and Contact Pages",
        "Guest / Client / Admin (public information)",
        "03-about.png",
        "The About and Contact pages provide company information and a way to reach Deen "
        "Associate. These pages do not require login.",
        steps=[
            "Use About to read company background and mission content.",
            "Use Contact to find phone, email, or inquiry options.",
        ],
    )
    add_figure(doc, 12, "Contact Page", "04-contact.png")

    add_heading(doc, "5. Administrator User Guide", level=1)
    add_paragraph(
        doc,
        "Administrators use the expanded navigation bar to run daily operations. The sections "
        "below follow the typical workflow: requests → confirmed bookings → payments → "
        "documents → staff → finance.",
    )

    add_heading(doc, "5.1 Booking Requests and Confirmed Bookings", level=2)
    add_paragraph(
        doc,
        "Requests lists incoming booking inquiries submitted by clients. After verification, "
        "an admin approves the request and creates a confirmed booking. The Bookings module "
        "then tracks the full financial lifecycle of that sale.",
    )

    add_screen_section(
        doc,
        13,
        "Confirmed Booking Detail and Installment Setup",
        "Admin only",
        "11-booking-detail-admin.png",
        "The booking detail page (example: BK-000007) is the control center for a confirmed "
        "sale. Summary cards show agreed sale price, booking amount received/remaining, "
        "and installment pool. The Regenerate Installment Plan section lets the admin "
        "negotiate final terms: sale price, discount, frequency, number of installments, "
        "start date, and optional possession amount.",
        steps=[
            "Open Bookings and select a confirmed booking.",
            "Review financial summary cards at the top of the page.",
            "Enter or adjust plan fields in Regenerate Installment Plan.",
            "Check the preview line (for example, 12 installments × amount).",
            "Click Regenerate Schedule to rebuild the installment table.",
            "Use Print Application Form or Cancel Booking when needed.",
        ],
        notes=(
            "The Payment Plan Active badge indicates an installment schedule is currently "
            "in force for the booking."
        ),
    )

    add_screen_section(
        doc,
        14,
        "Installment Schedule",
        "Admin only",
        "12-installment-schedule.png",
        "After a plan is generated, the Installment Schedule table lists each installment "
        "with sequence number, type, due date, amount, paid amount, remaining balance, "
        "status, and notes. The header summarizes total, paid, and remaining amounts.",
        steps=[
            "Scroll to Installment Schedule on the booking detail page.",
            "Review due dates and pending or paid statuses.",
            "Click Record Payment on the row you are collecting.",
        ],
    )

    add_screen_section(
        doc,
        15,
        "Record Installment Payment Dialog",
        "Admin only",
        "13-record-payment-modal.png",
        "The Record Installment Payment dialog captures collection details: amount (defaults "
        "to remaining balance), payment method (Cash, bank transfer, etc.), optional "
        "reference, payment date, and notes. Saving updates the schedule and creates a "
        "receipt in the finance ledger.",
        steps=[
            "Click Record Payment on an installment row.",
            "Confirm or edit the amount and payment method.",
            "Enter reference and payment date if applicable.",
            "Click Record Payment to post the transaction.",
            "Open the generated receipt to print for the customer.",
        ],
    )

    add_heading(doc, "5.2 Customers Module", level=2)
    add_paragraph(
        doc,
        "Customers stores registered client profiles linked to bookings. Administrators "
        "use this module to review contact details, view booking history, and maintain "
        "customer records. Access it from the Customers link in the admin navigation bar.",
    )
    add_bullets(
        doc,
        [
            "Search customers by name, email, or phone.",
            "Open a customer to inspect linked bookings and profile fields.",
            "Use customer records when creating manual bookings or verifying identity.",
        ],
    )

    add_heading(doc, "5.3 Employees Module", level=2)
    add_paragraph(
        doc,
        "Employees is an admin-only HR area with three tabs: Team, Attendance, and Salaries.",
    )

    add_screen_section(
        doc,
        16,
        "Employees — Team Tab",
        "Admin only",
        "16-employees-team.png",
        "The Team tab lists all employees with position, department, status, assignments, "
        "phone, and join date. Use Add Employee to onboard staff and View Details to open "
        "a full profile.",
        steps=[
            "Click Employees in the navigation bar.",
            "Ensure the Team tab is selected.",
            "Search or filter by status and department.",
            "Add or open an employee record as needed.",
        ],
    )

    add_screen_section(
        doc,
        17,
        "Employees — Attendance Tab",
        "Admin only",
        "17-employees-attendance.png",
        "The Attendance tab records daily presence. Select a date, mark each employee "
        "Present or Absent, optionally capture a reason, then Save a row or Save All.",
        steps=[
            "Switch to the Attendance tab.",
            "Choose the attendance date.",
            "Mark Present or Absent for each employee.",
            "Click Save on a row or Save All to store the sheet.",
        ],
    )

    add_screen_section(
        doc,
        18,
        "Employees — Salaries Tab",
        "Admin only",
        "18-employees-salaries.png",
        "The Salaries tab generates monthly payroll. Select the month, review each "
        "employee's salary, mark Paid when disbursed, and View Receipt for processed "
        "salary payments. Salary payments also appear in Finance as expense entries.",
        steps=[
            "Switch to the Salaries tab.",
            "Choose the payroll month.",
            "Confirm salary amounts (editable where permitted).",
            "Mark Paid and open the receipt after payment.",
        ],
    )

    add_heading(doc, "5.4 Finance Module", level=2)
    add_paragraph(
        doc,
        "Finance gives administrators a company-wide view of revenue, expenses, profit, "
        "outstanding receivables, and overdue installments. Revenue includes automatic "
        "entries from booking and installment payments plus manually entered income.",
    )

    add_screen_section(
        doc,
        19,
        "Finance Dashboard — Revenue View",
        "Admin only",
        "19-finance-revenue.png",
        "The Revenue view shows summary cards and a detailed table. Automatic Payment "
        "rows link to receipts (for example, RCP-000006). Manual Revenue rows can be "
        "edited or deleted.",
        steps=[
            "Click Finance in the navigation bar.",
            "Review Total Revenue, Expenses, Net Profit, Outstanding, and Overdue cards.",
            "Keep the Revenue tab selected to audit incoming money.",
            "Use Add Revenue for manual income not tied to a booking payment.",
        ],
    )

    add_screen_section(
        doc,
        20,
        "Add Manual Revenue Dialog",
        "Admin only",
        "20-add-manual-revenue.png",
        "Manual revenue records non-booking income such as transfer charges, documentation "
        "fees, parking charges, or other project income. Select project (or General), "
        "revenue type, amount, date, and optional reference/description.",
        steps=[
            "On the Finance page, click Add Revenue.",
            "Choose project and revenue type.",
            "Enter amount and transaction date.",
            "Save to post the entry to the revenue ledger.",
        ],
    )

    add_screen_section(
        doc,
        21,
        "Finance Dashboard — Expenses View",
        "Admin only",
        "21-finance-expenses.png",
        "The Expenses view lists operational costs such as salaries and materials. Each "
        "row shows date, project, category, amount, description, and reference. Admins can "
        "Add Expense or edit/delete manual expense rows.",
        steps=[
            "Switch the Finance toggle to Expenses.",
            "Review categories and amounts.",
            "Click Add Expense for new costs.",
            "Use Edit or Delete on manual expense entries when corrections are required.",
        ],
    )

    add_heading(doc, "5.5 Project and Media Management", level=2)
    add_paragraph(
        doc,
        "Administrators maintain project content through the same project detail pages "
        "clients use to browse, but with additional create/edit/upload capabilities "
        "available from admin workflows (project setup, unit inventory, and media uploads). "
        "The Overview, Units, and Media tabs shown in Section 4 also apply to admin "
        "maintenance—admins ensure descriptions, statuses, and gallery assets remain "
        "accurate for client-facing pages.",
    )

    add_heading(doc, "6. Quick Reference — Who Can Access What", level=1)
    rows = [
        ("Screen / Module", "Guest", "Client", "Admin"),
        ("Home, About, Contact", "Yes", "Yes", "Yes"),
        ("Projects & unit browsing", "Yes", "Yes", "Yes"),
        ("Sign Up / Log in", "Yes", "Yes", "Yes"),
        ("Submit booking request", "No", "Yes", "Yes"),
        ("My Projects", "No", "Yes", "No"),
        ("Requests", "No", "No", "Yes"),
        ("Bookings (admin list & detail)", "No", "No", "Yes"),
        ("Record installment payment", "No", "No", "Yes"),
        ("Customers", "No", "No", "Yes"),
        ("Employees", "No", "No", "Yes"),
        ("Finance", "No", "No", "Yes"),
        ("Application form & receipt (own booking)", "No", "Yes", "Yes (all)"),
    ]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, size=10, bold=(r == 0))

    doc.add_paragraph()
    add_heading(doc, "7. Tips and Troubleshooting", level=1)
    add_bullets(
        doc,
        [
            "If pages appear empty after login, confirm the backend API is running and the browser has network access.",
            "Clients who cannot see My Projects should verify the booking is confirmed and linked to their account.",
            "Installment totals must be regenerated after changing agreed price or discount on a booking.",
            "Use Print / Save as PDF for application forms and receipts rather than screenshots for official records.",
            "Finance outstanding and overdue cards update when installment payments are recorded against bookings.",
            "Log out from the profile menu when using a shared computer.",
        ],
    )

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
