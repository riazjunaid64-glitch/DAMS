"""Generate DAMS Sample Queries document in PDF and DOCX (thesis sample style)."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
OUT_PDF = ROOT / "docs" / "DAMS_Sample_Queries.pdf"
OUT_DOCX = ROOT / "docs" / "DAMS_Sample_Queries.docx"

DOC_DATE = "27th November, 2025"
AUTHOR = "Muhammad Junaid Riaz"
SYSTEM = "Deen Associate Management System"


@dataclass(frozen=True)
class SampleQuery:
    number: int
    name: str
    description: str
    sql: str
    service: str
    endpoint: str


def build_queries() -> list[SampleQuery]:
    return [
        SampleQuery(
            1,
            "User Login Verification",
            "Verify registered user credentials during login.",
            """SELECT u.UserId, u.FullName, u.Email, u.Password, r.Role_name
FROM Users u
INNER JOIN Roles r ON u.RoleId = r.RoleId
WHERE u.Email = @Email;""",
            "AuthService.Login()",
            "POST /api/Auth/login",
        ),
        SampleQuery(
            2,
            "Check Duplicate Email on Registration",
            "Prevent duplicate client registration using the same email address.",
            """SELECT UserId
FROM Users
WHERE LOWER(Email) = LOWER(@Email);""",
            "AuthService.RegisterAsync()",
            "POST /api/Auth/register",
        ),
        SampleQuery(
            3,
            "Insert New Client User",
            "Register a new client user account with hashed password.",
            """INSERT INTO Users (FullName, Email, Password, RoleId)
VALUES (@FullName, @Email, @HashedPassword, @ClientRoleId);""",
            "AuthService.RegisterAsync()",
            "POST /api/Auth/register",
        ),
        SampleQuery(
            4,
            "Retrieve All Projects",
            "Fetch all real-estate projects for public listing and admin dashboard.",
            """SELECT Id, ProjectName, Location, Description, StartingDate,
       ExpectedCompletionDate, Status, CreatedAt
FROM Projects
ORDER BY CreatedAt DESC;""",
            "ProjectService.GetAllProjectsAsync()",
            "GET /api/Project",
        ),
        SampleQuery(
            5,
            "Retrieve Project by ID",
            "Fetch details of a single project.",
            """SELECT Id, ProjectName, Location, Description, StartingDate,
       ExpectedCompletionDate, Status, CreatedAt
FROM Projects
WHERE Id = @ProjectId;""",
            "ProjectService.GetProjectByIdAsync()",
            "GET /api/Project/{id}",
        ),
        SampleQuery(
            6,
            "Retrieve Units by Project",
            "List all units belonging to a selected project.",
            """SELECT Id, ProjectId, UnitNumber, UnitType, FloorNumber, Size, Price, Status
FROM Units
WHERE ProjectId = @ProjectId
ORDER BY FloorNumber, UnitNumber;""",
            "UnitService.GetUnitsByProjectIdAsync()",
            "GET /api/Unit/project/{projectId}",
        ),
        SampleQuery(
            7,
            "Search and Filter Customers",
            "Admin searches customers by name, phone, CNIC, or email with optional filters.",
            """SELECT c.*, COUNT(b.Id) AS BookingsCount
FROM Customers c
LEFT JOIN Bookings b ON b.CustomerId = c.Id
WHERE (@Source IS NULL OR c.Source = @Source)
  AND (@Status IS NULL OR c.Status = @Status)
  AND (
        LOWER(c.FullName) LIKE '%' + LOWER(@SearchTerm) + '%'
     OR c.Phone LIKE '%' + @SearchTerm + '%'
     OR c.CNIC LIKE '%' + @SearchTerm + '%'
     OR LOWER(c.Email) LIKE '%' + LOWER(@SearchTerm) + '%'
      )
GROUP BY c.Id
ORDER BY c.CreatedAt DESC
OFFSET @Skip ROWS FETCH NEXT @PageSize ROWS ONLY;""",
            "CustomerService.GetCustomersAsync()",
            "GET /api/Customer?search=&source=&status=",
        ),
        SampleQuery(
            8,
            "Find Existing Customer (Deduplication)",
            "Locate an existing customer by CNIC, then phone, then email before creating a new record.",
            """SELECT TOP 1 Id FROM Customers WHERE CNIC = @CNIC;
-- if not found:
SELECT TOP 1 Id FROM Customers WHERE Phone = @Phone;
-- if not found:
SELECT TOP 1 Id FROM Customers WHERE Email = @Email;""",
            "CustomerService.FindOrCreateCustomerAsync()",
            "Used internally by BookingService / BookingRequestService",
        ),
        SampleQuery(
            9,
            "Insert Booking Request",
            "Client submits an online booking request for an available unit.",
            """INSERT INTO BookingRequests
    (UnitId, FullName, Phone, CNIC, Email, Address, Status, UserId, CreatedAt)
VALUES
    (@UnitId, @FullName, @Phone, @CNIC, @Email, @Address, 'Pending', @UserId, GETUTCDATE());""",
            "BookingRequestService.SubmitBookingRequestAsync()",
            "POST /api/BookingRequest",
        ),
        SampleQuery(
            10,
            "List Booking Requests (Admin)",
            "Admin views booking requests with project, unit, and status filters.",
            """SELECT br.*, u.UnitNumber, p.ProjectName
FROM BookingRequests br
INNER JOIN Units u ON br.UnitId = u.Id
INNER JOIN Projects p ON u.ProjectId = p.Id
WHERE (@Status IS NULL OR br.Status = @Status)
  AND (@ProjectId IS NULL OR u.ProjectId = @ProjectId)
ORDER BY br.CreatedAt DESC
OFFSET @Skip ROWS FETCH NEXT @PageSize ROWS ONLY;""",
            "BookingRequestService.GetBookingRequestsAsync()",
            "GET /api/BookingRequest?status=&projectId=&search=",
        ),
        SampleQuery(
            11,
            "Approve Booking Request",
            "Admin approves a pending request, creates customer/booking, and marks unit as booked.",
            """UPDATE BookingRequests
SET Status = 'Approved', ReviewedByUserId = @AdminUserId, ReviewedAt = GETUTCDATE()
WHERE Id = @BookingRequestId AND Status = 'Pending';

INSERT INTO Bookings (CustomerId, UnitId, BookingRequestId, Status, ...)
VALUES (@CustomerId, @UnitId, @BookingRequestId, 'AwaitingBookingAmount', ...);

UPDATE Units SET Status = 'Booked' WHERE Id = @UnitId;""",
            "BookingRequestService.ApproveBookingRequestAsync()",
            "POST /api/BookingRequest/{id}/approve",
        ),
        SampleQuery(
            12,
            "Create Walk-in Booking",
            "Admin creates a booking directly for an available unit.",
            """SELECT u.*, p.ProjectName
FROM Units u
INNER JOIN Projects p ON u.ProjectId = p.Id
WHERE u.Id = @UnitId AND u.Status = 'Available';

INSERT INTO Bookings
    (CustomerId, UnitId, Status, AgreedSalePrice, BookingAmountRequired, CreatedByUserId, CreatedAt)
VALUES
    (@CustomerId, @UnitId, 'AwaitingBookingAmount', @AgreedSalePrice, @BookingAmountRequired, @AdminUserId, GETUTCDATE());

UPDATE Units SET Status = 'Booked' WHERE Id = @UnitId;""",
            "BookingService.CreateBookingAsync()",
            "POST /api/Booking",
        ),
        SampleQuery(
            13,
            "List Bookings with Search and Pagination",
            "Admin retrieves bookings filtered by status, project, customer, or search term.",
            """SELECT b.*, c.FullName, u.UnitNumber, p.ProjectName
FROM Bookings b
INNER JOIN Customers c ON b.CustomerId = c.Id
INNER JOIN Units u ON b.UnitId = u.Id
INNER JOIN Projects p ON u.ProjectId = p.Id
WHERE (@Status IS NULL OR b.Status = @Status)
  AND (@ProjectId IS NULL OR u.ProjectId = @ProjectId)
  AND (
        LOWER(b.BookingReference) LIKE '%' + LOWER(@SearchTerm) + '%'
     OR LOWER(c.FullName) LIKE '%' + LOWER(@SearchTerm) + '%'
     OR LOWER(u.UnitNumber) LIKE '%' + LOWER(@SearchTerm) + '%'
      )
ORDER BY b.BookingDate DESC
OFFSET @Skip ROWS FETCH NEXT @PageSize ROWS ONLY;""",
            "BookingService.GetBookingsAsync()",
            "GET /api/Booking?status=&projectId=&search=",
        ),
        SampleQuery(
            14,
            "Client View Own Bookings (My Projects)",
            "Logged-in client views bookings linked to the same email as the JWT account.",
            """SELECT b.*, c.FullName, u.UnitNumber, p.ProjectName
FROM Bookings b
INNER JOIN Customers c ON b.CustomerId = c.Id
INNER JOIN Units u ON b.UnitId = u.Id
INNER JOIN Projects p ON u.ProjectId = p.Id
WHERE LOWER(c.Email) = LOWER(@JwtEmail)
  AND b.Status <> 'Cancelled'
ORDER BY b.BookingDate DESC;""",
            "BookingService.GetBookingsByCustomerEmailAsync()",
            "GET /api/MyProjects",
        ),
        SampleQuery(
            15,
            "Record Booking Amount Payment",
            "Admin records initial booking amount payment and updates booking totals.",
            """INSERT INTO Payments
    (BookingId, Type, Amount, PaymentMethod, Reference, PaidAt, ReceiptNumber, CreatedByUserId)
VALUES
    (@BookingId, 'BookingAmount', @Amount, @PaymentMethod, @Reference, @PaidAt, @ReceiptNumber, @AdminUserId);

UPDATE Bookings
SET BookingAmountReceived = BookingAmountReceived + @Amount,
    Status = CASE WHEN BookingAmountReceived + @Amount >= BookingAmountRequired
                  THEN 'PaymentPlanActive' ELSE Status END
WHERE Id = @BookingId;""",
            "BookingService.RecordBookingPaymentAsync()",
            "POST /api/Booking/{id}/payments",
        ),
        SampleQuery(
            16,
            "Retrieve Installment Schedule",
            "Fetch installment plan rows for a booking.",
            """SELECT i.*
FROM Installments i
INNER JOIN Bookings b ON i.BookingId = b.Id
WHERE b.Id = @BookingId
ORDER BY i.SequenceNumber;""",
            "InstallmentService.GetScheduleAsync()",
            "GET /api/Booking/{id}/installments",
        ),
        SampleQuery(
            17,
            "Record Installment Payment",
            "Admin records payment against a specific installment.",
            """INSERT INTO Payments
    (BookingId, InstallmentId, Type, Amount, PaymentMethod, Reference, PaidAt, ReceiptNumber)
VALUES
    (@BookingId, @InstallmentId, 'Installment', @Amount, @PaymentMethod, @Reference, @PaidAt, @ReceiptNumber);

UPDATE Installments
SET Status = CASE WHEN @RemainingAfterPayment <= 0 THEN 'Paid' ELSE 'PartiallyPaid' END
WHERE Id = @InstallmentId;""",
            "InstallmentService.RecordInstallmentPaymentAsync()",
            "POST /api/Booking/{id}/installments/{installmentId}/payment",
        ),
        SampleQuery(
            18,
            "Finance Dashboard — Automatic Revenue",
            "Aggregate booking and installment payments as automatic revenue.",
            """SELECT p.PaidAt, p.Amount, p.Type, p.ReceiptNumber,
       b.BookingReference, c.FullName, pr.ProjectName
FROM Payments p
INNER JOIN Bookings b ON p.BookingId = b.Id
INNER JOIN Customers c ON b.CustomerId = c.Id
INNER JOIN Units u ON b.UnitId = u.Id
LEFT JOIN Projects pr ON u.ProjectId = pr.Id
WHERE b.Status <> 'Cancelled'
  AND (@ProjectId IS NULL OR u.ProjectId = @ProjectId)
  AND (@FromDate IS NULL OR p.PaidAt >= @FromDate)
  AND (@ToDate IS NULL OR p.PaidAt < @ToDateExclusive);""",
            "FinanceService.GetDashboardAsync()",
            "GET /api/Finance/dashboard?projectId=&from=&to=",
        ),
        SampleQuery(
            19,
            "Finance Dashboard — Manual Revenue",
            "Retrieve manually entered project revenue records.",
            """SELECT mr.Id, mr.Date, mr.RevenueType, mr.Amount, mr.Reference, mr.Description,
       p.ProjectName
FROM ManualRevenues mr
LEFT JOIN Projects p ON mr.ProjectId = p.Id
WHERE (@ProjectId IS NULL OR mr.ProjectId = @ProjectId)
  AND (@FromDate IS NULL OR mr.Date >= @FromDate)
  AND (@ToDate IS NULL OR mr.Date < @ToDateExclusive);""",
            "FinanceService.GetDashboardAsync()",
            "GET /api/Finance/dashboard?projectId=&from=&to=",
        ),
        SampleQuery(
            20,
            "Finance Dashboard — Expenses",
            "Retrieve project and general expenses including salary expenses.",
            """SELECT e.Id, e.Date, e.Category, e.Amount, e.Vendor, e.Description,
       p.ProjectName
FROM Expenses e
LEFT JOIN Projects p ON e.ProjectId = p.Id
WHERE (@ProjectId IS NULL OR e.ProjectId = @ProjectId)
  AND (@FromDate IS NULL OR e.Date >= @FromDate)
  AND (@ToDate IS NULL OR e.Date < @ToDateExclusive);""",
            "FinanceService.GetDashboardAsync()",
            "GET /api/Finance/dashboard?projectId=&from=&to=",
        ),
        SampleQuery(
            21,
            "List Employees with Filters",
            "Admin retrieves employees filtered by department and status.",
            """SELECT Id, FullName, Department, Salary, Status, JoinDate
FROM Employees
WHERE (@Department IS NULL OR Department = @Department)
  AND (@Status IS NULL OR Status = @Status)
ORDER BY FullName;""",
            "EmployeeService.GetAllEmployeesAsync()",
            "GET /api/Employee?department=&status=",
        ),
        SampleQuery(
            22,
            "Record Employee Attendance",
            "Mark daily attendance for an employee on a selected date.",
            """SELECT Id FROM EmployeeAttendances
WHERE EmployeeId = @EmployeeId AND [Date] = @Date;

-- if not exists:
INSERT INTO EmployeeAttendances (EmployeeId, Date, Status, Notes, RecordedByUserId)
VALUES (@EmployeeId, @Date, @Status, @Notes, @AdminUserId);""",
            "EmployeeService.RecordAttendanceAsync()",
            "POST /api/Employee/{id}/attendance",
        ),
        SampleQuery(
            23,
            "Generate Employee Salary",
            "Create monthly salary record and linked expense entry.",
            """SELECT Id FROM EmployeeSalaries
WHERE EmployeeId = @EmployeeId AND PayMonth = @Month AND PayYear = @Year;

INSERT INTO Expenses (Category, Amount, Vendor, Date, ProjectId, CreatedByUserId)
VALUES ('Salary', @Amount, @EmployeeName, @PayDate, @ProjectId, @AdminUserId);

INSERT INTO EmployeeSalaries (EmployeeId, Amount, PayMonth, PayYear, ExpenseId, PayDate)
VALUES (@EmployeeId, @Amount, @Month, @Year, @ExpenseId, @PayDate);""",
            "EmployeeService.GenerateSalaryAsync()",
            "POST /api/Employee/{employeeId}/salary",
        ),
        SampleQuery(
            24,
            "Cancel Booking and Release Unit",
            "Admin cancels an active booking and returns the unit to Available status.",
            """UPDATE Bookings
SET Status = 'Cancelled', UpdatedAt = GETUTCDATE(), InternalNotes = @UpdatedNotes
WHERE Id = @BookingId AND Status NOT IN ('Cancelled', 'PossessionGiven', 'SaleCompleted');

UPDATE Units
SET Status = 'Available', UpdatedAt = GETUTCDATE()
WHERE Id = @UnitId;""",
            "BookingService.CancelBookingAsync()",
            "POST /api/Booking/{id}/cancel",
        ),
        SampleQuery(
            25,
            "Booking Request Statistics",
            "Count booking requests grouped by status for admin dashboard cards.",
            """SELECT Status, COUNT(*) AS RequestCount
FROM BookingRequests
GROUP BY Status;""",
            "BookingRequestService.GetBookingRequestStatsAsync()",
            "GET /api/BookingRequest/stats",
        ),
    ]


# ── DOCX helpers ──────────────────────────────────────────────────────────────

GREY = "D9D9D9"


def set_cell_shading(cell, fill: str = GREY) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def write_cell(cell, text: str, bold: bool = False, size: int = 10, mono: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Courier New" if mono else "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)


def add_cover_docx(doc: Document) -> None:
    for _ in range(6):
        doc.add_paragraph()

    def center(text: str, size: int = 14, bold: bool = False) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)

    center("Sample Queries", 18, True)
    center("FOR", 14, True)
    center(SYSTEM, 18, True)
    center("VERSION 1.0", 12)
    doc.add_paragraph()
    center("Prepared By", 12)
    center(AUTHOR, 14, True)
    doc.add_paragraph()
    center(DOC_DATE, 12)
    doc.add_page_break()


def build_docx(queries: list[SampleQuery]) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    add_cover_docx(doc)

    doc.add_heading("Revision History", level=1)
    rev = doc.add_table(rows=2, cols=4)
    rev.style = "Table Grid"
    headers = ["Version", "Description", "Author", "Date"]
    for i, h in enumerate(headers):
        write_cell(rev.rows[0].cells[i], h, bold=True)
        set_cell_shading(rev.rows[0].cells[i])
    write_cell(rev.rows[1].cells[0], "1.0")
    write_cell(
        rev.rows[1].cells[1],
        f"This document contains sample database queries for {SYSTEM} (DAMS).",
    )
    write_cell(rev.rows[1].cells[2], AUTHOR)
    write_cell(rev.rows[1].cells[3], DOC_DATE)
    doc.add_page_break()

    doc.add_heading("Sample Queries", level=1)
    intro = doc.add_paragraph()
    intro.add_run(
        "This section presents sample queries used by the Deen Associate Management System (DAMS). "
        "The application implements these operations using Entity Framework Core (LINQ) in the backend "
        "services. The SQL statements below show the equivalent database operations executed against "
        "Microsoft SQL Server."
    )

    doc.add_heading("Query Summary Table", level=2)
    summary = doc.add_table(rows=1, cols=4)
    summary.style = "Table Grid"
    for i, h in enumerate(["Query No.", "Query Name", "Service Method", "API Endpoint"]):
        write_cell(summary.rows[0].cells[i], h, bold=True)
        set_cell_shading(summary.rows[0].cells[i])
    for q in queries:
        row = summary.add_row().cells
        write_cell(row[0], str(q.number))
        write_cell(row[1], q.name)
        write_cell(row[2], q.service, size=9)
        write_cell(row[3], q.endpoint, size=9)

    doc.add_page_break()
    doc.add_heading("Detailed Sample Queries", level=2)

    for q in queries:
        doc.add_heading(f"Query {q.number}: {q.name}", level=3)
        p = doc.add_paragraph()
        p.add_run("Description: ").bold = True
        p.add_run(q.description)
        p = doc.add_paragraph()
        p.add_run("Service / Endpoint: ").bold = True
        p.add_run(f"{q.service}  |  {q.endpoint}")
        p = doc.add_paragraph()
        p.add_run("Equivalent SQL Query:").bold = True
        add_code_block_docx(doc, q.sql)
        doc.add_paragraph()

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(f"Created {OUT_DOCX}")


def add_code_block_docx(doc: Document, code: str) -> None:
    for line in code.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.15)
        run = p.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9)


# ── PDF helpers ───────────────────────────────────────────────────────────────

def esc(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def make_para(text: str, font: str = "Times-Roman", size: int = 10, bold: bool = False) -> Paragraph:
    style = ParagraphStyle(
        "Cell",
        fontName="Times-Bold" if bold else font,
        fontSize=size,
        leading=size + 3,
        alignment=TA_LEFT,
        wordWrap="CJK",
    )
    return Paragraph(esc(text), style)


def make_sql_para(sql: str) -> Paragraph:
    style = ParagraphStyle(
        "SQL",
        fontName="Courier",
        fontSize=8,
        leading=10,
        alignment=TA_LEFT,
        wordWrap="CJK",
    )
    return Paragraph(esc(sql).replace("\n", "<br/>"), style)


def build_wrapped_table(rows: list[list], col_widths: list[float], header: bool = True) -> Table:
    data = []
    for r_idx, row in enumerate(rows):
        if isinstance(row[0], Paragraph):
            data.append(row)
        else:
            data.append(
                [
                    make_para(str(cell), bold=(header and r_idx == 0 and isinstance(cell, str)))
                    if not isinstance(cell, Paragraph)
                    else cell
                    for cell in row
                ]
            )
    table = Table(data, colWidths=col_widths, repeatRows=1 if header else 0)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.whitesmoke]),
            ]
        )
    )
    return table


def build_pdf(queries: list[SampleQuery]) -> None:
    OUT_PDF.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )
    styles = getSampleStyleSheet()
    cover_title = ParagraphStyle("CT", fontName="Times-Bold", fontSize=18, leading=24, alignment=TA_CENTER)
    cover_sub = ParagraphStyle("CS", fontName="Times-Bold", fontSize=14, leading=18, alignment=TA_CENTER)
    cover_body = ParagraphStyle("CB", fontName="Times-Roman", fontSize=12, leading=16, alignment=TA_CENTER)
    heading = ParagraphStyle("H", fontName="Times-Bold", fontSize=13, leading=16, spaceBefore=10, spaceAfter=8)
    subheading = ParagraphStyle("SH", fontName="Times-Bold", fontSize=12, leading=15, spaceBefore=8, spaceAfter=6)
    body = ParagraphStyle("B", fontName="Times-Roman", fontSize=11, leading=15, alignment=TA_JUSTIFY, spaceAfter=6)

    page_width = letter[0] - doc.leftMargin - doc.rightMargin

    revision_table = build_wrapped_table(
        [
            ["Version", "Description", "Author", "Date"],
            [
                "1.0",
                f"This document contains sample database queries for {SYSTEM} (DAMS).",
                AUTHOR,
                DOC_DATE,
            ],
        ],
        [0.65 * inch, page_width - 2.95 * inch, 1.35 * inch, 0.95 * inch],
    )

    story = [
        Spacer(1, 1.4 * inch),
        Paragraph("SAMPLE QUERIES", cover_title),
        Spacer(1, 0.15 * inch),
        Paragraph("FOR", cover_sub),
        Spacer(1, 0.1 * inch),
        Paragraph(SYSTEM, cover_title),
        Spacer(1, 0.2 * inch),
        Paragraph("VERSION 1.0", cover_body),
        Spacer(1, 0.45 * inch),
        Paragraph("Prepared By", cover_body),
        Spacer(1, 0.1 * inch),
        Paragraph(AUTHOR, cover_sub),
        Spacer(1, 0.35 * inch),
        Paragraph(DOC_DATE, cover_body),
        PageBreak(),
        Paragraph("Revision History", heading),
        Spacer(1, 8),
        revision_table,
        PageBreak(),
        Paragraph("Sample Queries", heading),
        Paragraph(
            "This section presents sample queries used by the Deen Associate Management System (DAMS). "
            "The application implements these operations using Entity Framework Core (LINQ) in backend "
            "services. The SQL statements below show equivalent database operations against Microsoft SQL Server.",
            body,
        ),
        Paragraph("Query Summary Table", subheading),
    ]

    summary_rows: list[list] = [["Query No.", "Query Name", "Service Method", "API Endpoint"]]
    for q in queries:
        summary_rows.append([str(q.number), q.name, q.service, q.endpoint])
    story += [
        build_wrapped_table(
            summary_rows,
            [0.55 * inch, 1.55 * inch, 1.85 * inch, page_width - 3.95 * inch],
        ),
        PageBreak(),
        Paragraph("Detailed Sample Queries", subheading),
    ]

    for q in queries:
        story += [
            Paragraph(f"Query {q.number}: {q.name}", subheading),
            Paragraph(f"<b>Description:</b> {esc(q.description)}", body),
            Paragraph(f"<b>Service / Endpoint:</b> {esc(q.service)}  |  {esc(q.endpoint)}", body),
            Paragraph("<b>Equivalent SQL Query:</b>", body),
            build_wrapped_table(
                [[make_sql_para(q.sql)]],
                [page_width],
                header=False,
            ),
            Spacer(1, 10),
        ]

    doc.build(story)
    print(f"Created {OUT_PDF}")


def main() -> None:
    queries = build_queries()
    build_docx(queries)
    build_pdf(queries)


if __name__ == "__main__":
    main()
