"""Generate DAMS Project Code Word document for documentation."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "docs" / "DAMS_Project_Code.docx"


def add_heading(doc: Document, text: str, level: int = 2) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_code(doc: Document, code: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(code)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0, 0, 0)


SECTIONS: list[tuple[str, str]] = [
    (
        "4.2.1 AuthController",
        """using Microsoft.AspNetCore.Mvc;
using Microsoft.AspNetCore.Authorization;
using DAMS.Application.DTOs.Auth;
using DAMS.Application.Interfaces;
using System.Security.Claims;

namespace DAMS.Api.Controllers
{
    [ApiController]
    [Route("api/[controller]")]
    public class AuthController : ControllerBase
    {
        private readonly IAuthService _authService;

        public AuthController(IAuthService authService)
        {
            _authService = authService;
        }

        [HttpPost("register")]
        public async Task<IActionResult> Register(RegisterRequestDto request)
        {
            await _authService.RegisterAsync(request);
            return Ok(new { message = "Registration completed successfully." });
        }

        [HttpPost("login")]
        public IActionResult Login(LoginRequestDto request)
        {
            var tokens = _authService.Login(request);
            if (tokens == null)
                return Unauthorized("Invalid credentials");
            return Ok(tokens);
        }

        [Authorize]
        [HttpGet("profile")]
        public IActionResult Profile()
        {
            return Ok(new
            {
                userId = User.FindFirst(ClaimTypes.NameIdentifier)?.Value,
                email = User.FindFirst(ClaimTypes.Email)?.Value,
                role = User.FindFirst(ClaimTypes.Role)?.Value
            });
        }
    }
}""",
    ),
    (
        "4.2.2 ProjectController",
        """using DAMS.Application.DTOs.ProjectDtos;
using DAMS.Application.Interfaces;
using Microsoft.AspNetCore.Authorization;
using Microsoft.AspNetCore.Mvc;
using System.Security.Claims;

namespace DAMS.Api.Controllers
{
    [Route("api/[controller]")]
    [ApiController]
    public class ProjectController : ControllerBase
    {
        private readonly IProjectService _projectService;

        public ProjectController(IProjectService projectService)
        {
            _projectService = projectService;
        }

        [Authorize(Roles = "Admin")]
        [HttpPost]
        public async Task<IActionResult> CreateProject(CreateProjectDto dto)
        {
            var adminId = int.Parse(User.FindFirstValue(ClaimTypes.NameIdentifier)!);
            var result = await _projectService.CreateProjectAsync(dto, adminId);
            return Ok(result);
        }

        [Authorize(Roles = "Admin")]
        [HttpPut("{id}")]
        public async Task<IActionResult> UpdateProject(int id, UpdateProjectDto dto)
        {
            var result = await _projectService.UpdateProjectAsync(id, dto);
            return Ok(result);
        }

        [AllowAnonymous]
        [HttpGet]
        public async Task<IActionResult> GetAllProjects()
        {
            var result = await _projectService.GetAllProjectsAsync();
            return Ok(result);
        }

        [AllowAnonymous]
        [HttpGet("{id}")]
        public async Task<IActionResult> GetProjectById(int id)
        {
            var result = await _projectService.GetProjectByIdAsync(id);
            if (result == null) return NotFound();
            return Ok(result);
        }
    }
}""",
    ),
    (
        "4.2.3 UnitController",
        """using Microsoft.AspNetCore.Mvc;
using Microsoft.AspNetCore.Authorization;
using DAMS.Application.Interfaces;
using DAMS.Application.DTOs.UnitDtos;

namespace DAMS.Api.Controllers
{
    [Route("api/[controller]")]
    [ApiController]
    public class UnitController : ControllerBase
    {
        private readonly IUnitService _unitService;

        public UnitController(IUnitService unitService)
        {
            _unitService = unitService;
        }

        [Authorize(Roles = "Admin")]
        [HttpPost]
        public async Task<IActionResult> Create(CreateUnitDto dto)
        {
            var result = await _unitService.CreateUnitAsync(dto);
            return Ok(result);
        }

        [AllowAnonymous]
        [HttpGet("{id:int}")]
        public async Task<IActionResult> GetById(int id)
        {
            var result = await _unitService.GetUnitByIdAsync(id);
            if (result == null) return NotFound();
            return Ok(result);
        }

        [Authorize(Roles = "Admin")]
        [HttpPut("{id}")]
        public async Task<IActionResult> Update(int id, UpdateUnitDto dto)
        {
            var result = await _unitService.UpdateUnitAsync(id, dto);
            return Ok(result);
        }

        [Authorize(Roles = "Admin")]
        [HttpDelete("{id}")]
        public async Task<IActionResult> Delete(int id)
        {
            await _unitService.DeleteUnitAsync(id);
            return Ok("Unit deleted successfully");
        }
    }
}""",
    ),
    (
        "4.2.4 CustomerController",
        """using DAMS.Application.DTOs.CustomerDtos;
using DAMS.Application.Interfaces;
using Microsoft.AspNetCore.Authorization;
using Microsoft.AspNetCore.Mvc;
using System.Security.Claims;

namespace DAMS.Api.Controllers
{
    [Route("api/[controller]")]
    [ApiController]
    [Authorize(Roles = "Admin")]
    public class CustomerController : ControllerBase
    {
        private readonly ICustomerService _customerService;

        public CustomerController(ICustomerService customerService)
        {
            _customerService = customerService;
        }

        [HttpPost]
        public async Task<IActionResult> Create([FromBody] CreateCustomerDto dto)
        {
            var result = await _customerService.CreateCustomerAsync(dto, GetUserId());
            return Ok(result);
        }

        [HttpGet]
        public async Task<IActionResult> GetAll([FromQuery] string? search)
        {
            var filter = new CustomerFilterDto { SearchTerm = search };
            var result = await _customerService.GetCustomersAsync(filter);
            return Ok(result);
        }

        [HttpGet("{id:int}")]
        public async Task<IActionResult> GetById(int id)
        {
            var result = await _customerService.GetCustomerByIdAsync(id);
            if (result == null) return NotFound();
            return Ok(result);
        }

        [HttpPut("{id:int}")]
        public async Task<IActionResult> Update(int id, [FromBody] UpdateCustomerDto dto)
        {
            var result = await _customerService.UpdateCustomerAsync(id, dto);
            return Ok(result);
        }

        private int GetUserId() => int.Parse(User.FindFirstValue(ClaimTypes.NameIdentifier)!);
    }
}""",
    ),
    (
        "4.2.5 BookingRequestController",
        """using DAMS.Application.DTOs.BookingRequestDtos;
using DAMS.Application.Interfaces;
using Microsoft.AspNetCore.Authorization;
using Microsoft.AspNetCore.Mvc;
using System.Security.Claims;

namespace DAMS.Api.Controllers
{
    [Route("api/[controller]")]
    [ApiController]
    public class BookingRequestController : ControllerBase
    {
        private readonly IBookingRequestService _bookingRequestService;

        [HttpPost]
        public async Task<IActionResult> CreateBookingRequest([FromBody] CreateBookingRequestDto dto)
        {
            var result = await _bookingRequestService.CreateBookingRequestAsync(dto, null);
            return Ok(result);
        }

        [HttpGet]
        [Authorize(Roles = "Admin")]
        public async Task<IActionResult> GetBookingRequests()
        {
            var result = await _bookingRequestService.GetBookingRequestsAsync(new BookingRequestFilterDto());
            return Ok(result);
        }

        [HttpPost("{id:int}/approve")]
        [Authorize(Roles = "Admin")]
        public async Task<IActionResult> ApproveBookingRequest(int id)
        {
            var adminUserId = int.Parse(User.FindFirstValue(ClaimTypes.NameIdentifier)!);
            var result = await _bookingRequestService.ApproveBookingRequestAsync(id, adminUserId);
            return Ok(result);
        }

        [HttpPost("{id:int}/reject")]
        [Authorize(Roles = "Admin")]
        public async Task<IActionResult> RejectBookingRequest(int id, [FromBody] RejectBookingRequestDto? dto)
        {
            var adminUserId = int.Parse(User.FindFirstValue(ClaimTypes.NameIdentifier)!);
            var result = await _bookingRequestService.RejectBookingRequestAsync(id, adminUserId, dto?.RejectionReason);
            return Ok(result);
        }
    }
}""",
    ),
    (
        "4.2.6 BookingController",
        """using DAMS.Application.DTOs.BookingDtos;
using DAMS.Application.Interfaces;
using Microsoft.AspNetCore.Authorization;
using Microsoft.AspNetCore.Mvc;
using System.Security.Claims;

namespace DAMS.Api.Controllers
{
    [Route("api/[controller]")]
    [ApiController]
    [Authorize(Roles = "Admin")]
    public class BookingController : ControllerBase
    {
        private readonly IBookingService _bookingService;

        [HttpPost]
        public async Task<IActionResult> Create([FromBody] CreateBookingDto dto)
        {
            var result = await _bookingService.CreateBookingAsync(dto, GetUserId());
            return Ok(result);
        }

        [HttpGet("{id:int}")]
        public async Task<IActionResult> GetById(int id)
        {
            var result = await _bookingService.GetBookingByIdAsync(id);
            if (result == null) return NotFound();
            return Ok(result);
        }

        [HttpPost("{id:int}/cancel")]
        public async Task<IActionResult> Cancel(int id, [FromBody] CancelBookingDto? dto)
        {
            var result = await _bookingService.CancelBookingAsync(id, dto?.Reason, GetUserId());
            return Ok(result);
        }

        [HttpPost("{id:int}/booking-amount-payment")]
        public async Task<IActionResult> RecordBookingAmountPayment(int id, [FromBody] RecordBookingAmountPaymentDto dto)
        {
            var result = await _bookingService.RecordBookingAmountPaymentAsync(id, dto, GetUserId());
            return Ok(result);
        }

        private int GetUserId() => int.Parse(User.FindFirstValue(ClaimTypes.NameIdentifier)!);
    }
}""",
    ),
    (
        "4.2.7 EmployeeController",
        """using DAMS.Application.DTOs.EmployeeDtos;
using DAMS.Application.Interfaces;
using Microsoft.AspNetCore.Authorization;
using Microsoft.AspNetCore.Mvc;

namespace DAMS.Api.Controllers
{
    [Route("api/[controller]")]
    [ApiController]
    [Authorize(Roles = "Admin")]
    public class EmployeeController : ControllerBase
    {
        private readonly IEmployeeService _employeeService;

        [HttpPost]
        public async Task<IActionResult> Create([FromBody] CreateEmployeeDto dto)
        {
            var result = await _employeeService.CreateEmployeeAsync(dto);
            return Ok(result);
        }

        [HttpGet]
        public async Task<IActionResult> GetAll()
        {
            var result = await _employeeService.GetAllEmployeesAsync(null, null);
            return Ok(result);
        }

        [HttpPost("{employeeId:int}/attendance")]
        public async Task<IActionResult> RecordAttendance(int employeeId, [FromBody] RecordAttendanceDto dto)
        {
            var result = await _employeeService.RecordAttendanceAsync(employeeId, dto);
            return Ok(result);
        }

        [HttpPost("{employeeId:int}/salary")]
        public async Task<IActionResult> GenerateSalary(int employeeId, [FromBody] GenerateSalaryDto dto)
        {
            var result = await _employeeService.GenerateSalaryAsync(employeeId, dto, null);
            return Ok(result);
        }
    }
}""",
    ),
    (
        "4.2.8 FinanceController",
        """using DAMS.Application.DTOs.ExpenseDtos;
using DAMS.Application.DTOs.FinanceDtos;
using DAMS.Application.Interfaces;
using Microsoft.AspNetCore.Authorization;
using Microsoft.AspNetCore.Mvc;
using System.Security.Claims;

namespace DAMS.Api.Controllers
{
    [Route("api/[controller]")]
    [ApiController]
    [Authorize(Roles = "Admin")]
    public class FinanceController : ControllerBase
    {
        private readonly IFinanceService _financeService;

        [HttpGet("dashboard")]
        public async Task<IActionResult> GetDashboard()
        {
            var result = await _financeService.GetDashboardAsync(null, null, null);
            return Ok(result);
        }

        [HttpPost("expenses")]
        public async Task<IActionResult> CreateExpense([FromBody] CreateExpenseDto dto)
        {
            var result = await _financeService.CreateExpenseAsync(dto, GetUserId());
            return Ok(result);
        }

        [HttpPost("revenue")]
        public async Task<IActionResult> CreateRevenue([FromBody] CreateManualRevenueDto dto)
        {
            var result = await _financeService.CreateManualRevenueAsync(dto, GetUserId());
            return Ok(result);
        }

        private int? GetUserId()
        {
            var claim = User.FindFirstValue(ClaimTypes.NameIdentifier);
            return int.TryParse(claim, out var id) ? id : null;
        }
    }
}""",
    ),
    (
        "4.2.9 AuthService",
        """using DAMS.Application.DTOs.Auth;
using DAMS.Application.Interfaces;
using DAMS.Domain.Entities;
using DAMS.Infrastructure.Data;
using Microsoft.EntityFrameworkCore;

namespace DAMS.Application.Services
{
    public class AuthService : IAuthService
    {
        private readonly AppDbContext _context;
        private readonly ITokenService _tokenService;

        public async Task RegisterAsync(RegisterRequestDto request)
        {
            var normalizedEmail = request.Email.Trim().ToLowerInvariant();
            var existingUser = await _context.Users
                .FirstOrDefaultAsync(u => u.Email.ToLower() == normalizedEmail);
            if (existingUser != null)
                throw new Exception("Email already exists");

            var clientRole = await _context.Roles
                .FirstOrDefaultAsync(r => r.Role_name.ToLower() == "client");

            var user = new User
            {
                FullName = request.FullName,
                Email = normalizedEmail,
                Password = BCrypt.Net.BCrypt.HashPassword(request.Password),
                RoleId = clientRole!.RoleId
            };

            _context.Users.Add(user);
            await _context.SaveChangesAsync();
        }

        public AuthResponseDto? Login(LoginRequestDto request)
        {
            var user = _context.Users.FirstOrDefault(u => u.Email == request.Email);
            if (user == null) return null;

            var isValid = BCrypt.Net.BCrypt.Verify(request.Password, user.Password);
            if (!isValid) return null;

            var role = _context.Roles.First(r => r.RoleId == user.RoleId);
            var accessToken = _tokenService.GenerateAccessToken(user, role.Role_name);
            var refreshToken = _tokenService.GenerateRefreshToken();

            user.RefreshToken = refreshToken;
            user.RefreshTokenExpiresAt = DateTime.UtcNow.AddDays(15);
            _context.SaveChanges();

            return new AuthResponseDto
            {
                AccessToken = accessToken,
                RefreshToken = refreshToken,
                ExpiresInMinutes = 15
            };
        }
    }
}""",
    ),
    (
        "4.2.10 BookingService",
        """using DAMS.Application.DTOs.BookingDtos;
using DAMS.Application.Interfaces;
using DAMS.Domain.Entities;
using DAMS.Domain.Enums;
using DAMS.Infrastructure.Data;
using Microsoft.EntityFrameworkCore;

namespace DAMS.Application.Services
{
    public class BookingService : IBookingService
    {
        private readonly AppDbContext _context;

        public async Task<BookingResponseDto> CreateBookingAsync(CreateBookingDto dto, int adminUserId)
        {
            var unit = await _context.Units.FirstOrDefaultAsync(u => u.Id == dto.UnitId);
            if (unit == null) throw new InvalidOperationException("Unit not found.");
            if (unit.Status != UnitStatus.Available)
                throw new InvalidOperationException("This unit is not available for booking.");

            var booking = new Booking
            {
                CustomerId = dto.CustomerId!.Value,
                UnitId = unit.Id,
                Status = BookingStatus.AwaitingBookingAmount,
                ListPrice = unit.Price,
                AgreedSalePrice = dto.AgreedSalePrice ?? unit.Price,
                BookingAmountRequired = dto.BookingAmountRequired ?? 0m,
                CreatedByUserId = adminUserId,
                CreatedAt = DateTime.UtcNow
            };

            _context.Bookings.Add(booking);
            unit.Status = UnitStatus.Booked;
            await _context.SaveChangesAsync();
            return await GetResponseAsync(booking.Id);
        }

        public async Task<BookingResponseDto> CancelBookingAsync(int id, string? reason, int adminUserId)
        {
            var booking = await _context.Bookings.Include(b => b.Unit)
                .FirstOrDefaultAsync(b => b.Id == id);
            if (booking == null) throw new InvalidOperationException("Booking not found.");

            booking.Status = BookingStatus.Cancelled;
            booking.Unit.Status = UnitStatus.Available;
            await _context.SaveChangesAsync();
            return await GetResponseAsync(booking.Id);
        }

        public async Task<BookingResponseDto> RecordBookingAmountPaymentAsync(
            int bookingId, RecordBookingAmountPaymentDto dto, int adminUserId)
        {
            var booking = await _context.Bookings.FirstOrDefaultAsync(b => b.Id == bookingId);
            if (booking == null) throw new InvalidOperationException("Booking not found.");

            booking.BookingAmountReceived += dto.Amount;
            if (booking.BookingAmountReceived >= booking.BookingAmountRequired)
                booking.Status = BookingStatus.PaymentPlanActive;

            await _context.SaveChangesAsync();
            return await GetResponseAsync(booking.Id);
        }
    }
}""",
    ),
]


def build_document() -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    add_heading(doc, "Project Code", level=1)

    for title, code in SECTIONS:
        add_heading(doc, f"{title}:", level=2)
        add_code(doc, code)
        doc.add_paragraph("")

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(f"Created {OUT_DOCX}")


if __name__ == "__main__":
    build_document()
