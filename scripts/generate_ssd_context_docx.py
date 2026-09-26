"""Generate complete DAMS SSD context Word document for external diagram tools."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "docs" / "DAMS_SSD_Context.docx"
GREY = "D9D9D9"


def set_run(run, size: int = 11, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_title(doc: Document, text: str, size: int = 18) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=size, bold=True)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run(r, size=16 if level == 1 else 14 if level == 2 else 12, bold=True)


def add_para(doc: Document, text: str, size: int = 11) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_run(r, size=size)


def add_bullets(doc: Document, items: list[str], size: int = 11) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        set_run(r, size=size)


def shade_cell(cell, fill: str = GREY) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], col_widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run(run, size=10, bold=True)
        shade_cell(cell)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run(run, size=9)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def add_entity_block(doc: Document, name: str, attrs: list[str], relations: list[str], notes: str = "") -> None:
    add_heading(doc, name, level=3)
    add_para(doc, "Attributes:")
    add_bullets(doc, attrs, size=10)
    add_para(doc, "Relationships:")
    add_bullets(doc, relations, size=10)
    if notes:
        add_para(doc, f"Notes: {notes}", size=10)


def add_ssd_sequence(doc: Document, title: str, actors: str, preconditions: str, steps: list[str], postconditions: str = "") -> None:
    add_heading(doc, title, level=2)
    add_para(doc, f"Primary actors: {actors}")
    add_para(doc, f"Preconditions: {preconditions}")
    add_para(doc, "Interaction sequence (for System Sequence Diagram):")
    for i, step in enumerate(steps, 1):
        add_para(doc, f"{i}. {step}", size=10)
    if postconditions:
        add_para(doc, f"Postconditions: {postconditions}")


def build() -> Path:
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.9)
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)

    add_title(doc, "DAMS — Complete System Context for SSD Generation")
    add_title(doc, "Deen Associate Management System", size=14)
    doc.add_paragraph()

    add_para(
        doc,
        "Purpose of this document: Provide end-to-end context about the Deen Associate "
        "Management System (DAMS) so an external tool (e.g. Claude) can create accurate "
        "System Sequence Diagrams (SSDs). This file describes architecture, all entities, "
        "roles, API endpoints, services, frontend routes, state transitions, and step-by-step "
        "interaction sequences for every major use case.",
    )

    # 1. Overview
    add_heading(doc, "1. Project Overview")
    add_para(
        doc,
        "DAMS is a real estate management web application for Deen Associate. It supports "
        "public project/unit browsing, online booking requests, admin booking confirmation, "
        "installment payment plans, printable application forms and receipts, customer CRM, "
        "employee HR (attendance, tasks, salaries), and company finance (revenue/expenses).",
    )
    add_table(
        doc,
        ["Item", "Detail"],
        [
            ["System name", "Deen Associate Management System (DAMS)"],
            ["Organization", "Deen Associate / Seven Ventures (project branding e.g. Floria Heights)"],
            ["Frontend", "React 18 + TypeScript + Vite + Tailwind CSS"],
            ["Backend", ".NET 8 Web API (C#)"],
            ["Database", "Microsoft SQL Server via Entity Framework Core"],
            ["Authentication", "JWT Bearer tokens + refresh tokens (BCrypt passwords)"],
            ["File storage", "Local wwwroot/uploads for project and unit media"],
            ["Architecture", "3-tier: Presentation (React) → Business Logic (.NET Services) → Data (EF Core / SQL Server)"],
        ],
        [1.6, 4.8],
    )

    # 2. Architecture
    add_heading(doc, "2. System Architecture")
    add_para(
        doc,
        "The frontend React SPA calls REST API endpoints on the .NET backend. Controllers "
        "delegate to Application-layer services. Services use AppDbContext (EF Core) directly — "
        "there is no separate repository layer. JWT middleware validates tokens before "
        "authorized controller actions.",
    )
    add_bullets(
        doc,
        [
            "Presentation Layer: frontend/src — pages, components, api client",
            "Business Logic Layer: BACKEND/DAMS.Application/Services — AuthService, BookingService, InstallmentService, etc.",
            "API Layer: BACKEND/DAMS.Api/Controllers — 9 controllers, ~80+ endpoints",
            "Domain Layer: BACKEND/DAMS.Domain — 17 entities, 16 enums",
            "Infrastructure Layer: BACKEND/DAMS.Infrastructure — AppDbContext, migrations",
        ],
    )

    # 3. Actors
    add_heading(doc, "3. Actors (for SSD Lifelines)")
    add_table(
        doc,
        ["Actor", "Description", "Typical SSD participation"],
        [
            ["Guest / Visitor", "Unauthenticated user browsing public pages", "Browse projects, submit booking request"],
            ["Client", "Registered customer (Role=Client)", "Login, My Projects, view receipts"],
            ["Admin", "Staff user (Role=Admin)", "All back-office operations"],
            ["Frontend UI", "React pages/modals (optional lifeline)", "Collects input, calls API, renders response"],
            ["AuthController", "API controller", "Register, login, refresh, profile"],
            ["BookingRequestController", "API controller", "Booking request CRUD and approve/reject"],
            ["BookingController", "API controller", "Confirmed bookings, payments, installments"],
            ["MyProjectsController", "API controller", "Client read-only booking access"],
            ["CustomerController", "API controller", "Customer CRM"],
            ["ProjectController / UnitController", "API controllers", "Projects, units, media"],
            ["EmployeeController", "API controller", "HR: employees, attendance, tasks, salary"],
            ["FinanceController", "API controller", "Finance dashboard, manual revenue, expenses"],
            ["AuthService", "Business service", "Register, login, token refresh"],
            ["BookingRequestService", "Business service", "Request workflow, approve/reject"],
            ["BookingService", "Business service", "Booking lifecycle, booking-amount payments, receipts"],
            ["InstallmentService", "Business service", "Generate schedule, record installment payments"],
            ["CustomerService", "Business service", "Customer CRUD, find-or-create deduplication"],
            ["EmployeeService", "Business service", "Attendance, tasks, salary + linked expense"],
            ["FinanceService", "Business service", "Dashboard aggregation, revenue/expense CRUD"],
            ["Database (SQL Server)", "Persistence", "All entity reads/writes via AppDbContext"],
        ],
        [1.3, 2.2, 2.3],
    )

    # 4. Roles
    add_heading(doc, "4. User Roles and Access Control")
    add_table(
        doc,
        ["Feature", "Guest", "Client", "Admin"],
        [
            ["Browse projects/units/media", "Yes", "Yes", "Yes"],
            ["Register / Login", "Yes", "Yes", "Yes"],
            ["Submit booking request", "Yes (no auth required)", "Yes", "Yes"],
            ["My Projects (own bookings)", "No", "Yes", "No"],
            ["View own receipts", "No", "Yes", "Yes (all)"],
            ["Booking Requests queue", "No", "No", "Yes"],
            ["Confirmed bookings management", "No", "No", "Yes"],
            ["Record payments / generate plans", "No", "No", "Yes"],
            ["Customers module", "No", "No", "Yes"],
            ["Employees module", "No", "No", "Yes"],
            ["Finance module", "No", "No", "Yes"],
            ["Create/edit projects/units/media", "No", "No", "Yes"],
        ],
        [2.0, 0.7, 0.7, 0.7],
    )
    add_para(
        doc,
        "Client access to bookings is matched by Customer.Email = JWT email claim (not Customer.UserId). "
        "JWT access token expires in 15 minutes; refresh token lasts 15 days. Roles seeded: Admin (RoleId=1), Client (RoleId=2).",
    )

    # 5. Entities
    add_heading(doc, "5. Domain Entities (All 17)")
    add_para(doc, "Every entity below exists in BACKEND/DAMS.Domain/Entities/ and is mapped in AppDbContext.")

    entities = [
        ("5.1 Role", ["RoleId (PK)", "Role_name (Admin, Client)"], ["1 → many User"], "Seeded at startup."),
        ("5.2 User", ["UserId (PK)", "RoleId (FK)", "FullName", "Email", "Password (BCrypt hash)", "RefreshToken", "RefreshTokenExpiresAt"], ["→ Role", "Optional link from Customer.UserId"], "Authentication account."),
        ("5.3 Customer", ["Id (PK)", "FullName", "FatherName", "Phone", "CNIC", "Email", "Address", "DateOfBirth", "Nationality", "Occupation", "Whatsapp", "Source (enum)", "SourceNotes", "Status (enum)", "UserId (FK, optional)", "Notes", "CreatedByUserId", "CreatedAt", "UpdatedAt"], ["→ User (optional)", "1 → many Booking"], "CRM record; may exist without login."),
        ("5.4 Project", ["Id (PK)", "ProjectName (unique)", "Location", "Description", "StartingDate", "ExpectedCompletionDate", "Status (enum)", "CreatedById", "CreatedAt", "UpdatedAt"], ["1 → many Unit", "1 → many ProjectMedia"], "Real estate development."),
        ("5.5 Unit", ["Id (PK)", "ProjectId (FK)", "UnitNumber", "UnitType", "FloorNumber", "Size", "Price", "Status (enum)", "CreatedAt", "UpdatedAt"], ["→ Project", "1 → many UnitMedia", "1 → many Booking"], "Individual property unit."),
        ("5.6 ProjectMedia", ["Id (PK)", "ProjectId (FK)", "MediaUrl", "MediaType", "Category (enum)", "IsCover", "DisplayOrder", "AltText", "Description", "FileSize", "Width", "Height", "OriginalFileName", "MimeType", "UploadedAt", "UpdatedAt"], ["→ Project (cascade delete)"], "Project gallery files."),
        ("5.7 UnitMedia", ["Id (PK)", "UnitId (FK)", "MediaUrl", "MediaType", "Category", "IsCover", "DisplayOrder", "AltText", "Description", "FileSize", "Width", "Height", "OriginalFileName", "MimeType", "UploadedAt", "UpdatedAt"], ["→ Unit (cascade delete)"], "Unit gallery files."),
        ("5.8 BookingRequest", ["Id (PK)", "UnitId (FK)", "UserId (FK, optional)", "FullName", "Phone", "Email", "CNIC", "Address", "Notes", "Status (Pending/Approved/Rejected)", "RequestedAt", "ReviewedAt", "ReviewedByUserId", "RejectionReason", "CustomerId (FK, set on approve)", "CreatedAt", "UpdatedAt"], ["→ Unit", "→ User (requester)", "→ User (reviewer)", "→ Customer"], "Online inquiry before confirmed booking."),
        ("5.9 Booking", ["Id (PK)", "BookingReference (BK-000123, unique)", "CustomerId (FK)", "UnitId (FK)", "BookingRequestId (FK, optional)", "Source", "AssignedSalesUserId", "Status (enum)", "ListPrice", "AgreedSalePrice", "DiscountAmount", "DiscountReason", "BookingAmountRequired", "BookingAmountReceived", "TotalInstallmentAmount", "BookingDate", "BookingAmountDueDate", "BookingAmountConfirmedDate", "InstallmentPlanStartDate", "InstallmentFrequency", "NumberOfInstallments", "PossessionAmount", "PossessionDueDate", "InstallmentPlanGeneratedAt", "PossessionDate", "CompletionDate", "CustomerNotes", "InternalNotes", "Application form snapshot fields (SerialNo, ApartmentCategory, Tower, IsCorner, PricePerSft, etc.)", "Next-of-kin fields", "CreatedByUserId", "CreatedAt", "UpdatedAt"], ["→ Customer", "→ Unit", "→ BookingRequest", "1 → many Installment", "1 → many Payment"], "Confirmed sale record."),
        ("5.10 Installment", ["Id (PK)", "BookingId (FK)", "SequenceNumber (0=Possession, 1..N=Regular)", "DueDate", "Type (Regular/Possession)", "Amount", "Status (Pending/Paid/PartiallyPaid/Overdue)", "PaidAt", "Notes"], ["→ Booking", "1 → many Payment", "Unique (BookingId, SequenceNumber)"], "Scheduled payment row."),
        ("5.11 Payment", ["Id (PK)", "BookingId (FK)", "InstallmentId (FK, null for booking amount)", "Type (BookingAmount/Installment)", "Amount", "PaymentMethod (Cash/BankTransfer/Cheque/Online)", "PaymentReference", "ReceiptNumber (RCP-000001, unique)", "Notes", "RecordedByUserId", "PaidAt", "CreatedAt"], ["→ Booking", "→ Installment (optional)"], "Financial transaction; auto revenue source."),
        ("5.12 Employee", ["Id (PK)", "FullName", "JobTitle", "Department", "Phone", "Email", "Address", "Salary (monthly base)", "JoinDate", "Status (enum)", "CreatedAt", "UpdatedAt"], ["1 → many EmployeeAttendance", "1 → many EmployeeTask", "1 → many EmployeeSalary"], "Staff record."),
        ("5.13 EmployeeAttendance", ["Id (PK)", "EmployeeId (FK)", "Date", "Status (Present/Absent/Late/HalfDay/Leave)", "CheckInTime", "CheckOutTime", "Notes", "CreatedAt"], ["→ Employee", "Unique (EmployeeId, Date)"], "Daily attendance."),
        ("5.14 EmployeeTask", ["Id (PK)", "EmployeeId (FK)", "ProjectId (FK, optional)", "Title", "Description", "Priority", "Status", "DueDate", "CompletedAt", "CreatedAt", "UpdatedAt"], ["→ Employee", "→ Project (optional)"], "Task assignment."),
        ("5.15 EmployeeSalary", ["Id (PK)", "EmployeeId (FK)", "Amount", "PayDate", "PayMonth", "PayYear", "ProjectId (FK, optional)", "ProjectName", "ExpenseId (FK, auto-created)", "Notes", "CreatedByUserId", "CreatedAt"], ["→ Employee", "→ Project", "→ Expense"], "Payroll record linked to finance expense."),
        ("5.16 Expense", ["Id (PK)", "ProjectId (FK, optional)", "Amount", "Category (e.g. Salary, Material)", "Description", "Vendor", "Date", "CreatedByUserId", "CreatedAt"], ["→ Project (optional)"], "Company expense; salary auto-created."),
        ("5.17 ManualRevenue", ["Id (PK)", "ProjectId (FK, optional)", "Amount", "RevenueType", "Description", "Reference", "Date", "CreatedByUserId", "CreatedAt"], ["→ Project (optional)"], "Non-booking income."),
    ]
    for name, attrs, rels, notes in entities:
        add_entity_block(doc, name, attrs, rels, notes)

    # 6. Enums
    add_heading(doc, "6. Enumerations")
    add_table(
        doc,
        ["Enum", "Values"],
        [
            ["UnitStatus", "Available, Booked, Reserved, Sold, PendingReview, OnPaymentPlan"],
            ["BookingStatus", "AwaitingBookingAmount, PaymentPlanActive, PossessionGiven, SaleCompleted, Cancelled"],
            ["BookingRequestStatus", "Pending, Approved, Rejected"],
            ["CustomerSource", "Website, WalkIn, Phone, Referral, Other"],
            ["CustomerStatus", "Active, Inactive, Blocked"],
            ["PaymentType", "BookingAmount, Installment"],
            ["PaymentMethod", "Cash, BankTransfer, Cheque, Online"],
            ["InstallmentType", "Regular, Possession"],
            ["InstallmentStatus", "Pending, Paid, Overdue, PartiallyPaid"],
            ["InstallmentFrequency", "Monthly, Quarterly, HalfYearly, Yearly"],
            ["ProjectStatus", "Planning, Ongoing, Completed, Cancelled, Archived"],
            ["MediaCategory", "Gallery, Thumbnail, FloorPlan, Brochure, ConstructionProgress, Interior, Exterior, Document, Video"],
            ["EmployeeStatus", "Active, Inactive, OnLeave, Terminated"],
            ["AttendanceStatus", "Present, Absent, Late, HalfDay, Leave"],
            ["EmployeeTaskStatus", "Pending, InProgress, Completed, Cancelled"],
            ["TaskPriority", "Low, Medium, High, Urgent"],
        ],
        [1.8, 4.0],
    )

    # 7. State transitions
    add_heading(doc, "7. State Transitions (Critical for SSDs)")
    add_heading(doc, "7.1 Unit Status Lifecycle", level=2)
    add_para(doc, "Available → PendingReview (booking request submitted) → Reserved (booking confirmed) → OnPaymentPlan (booking amount fully paid) → Available (booking cancelled/rejected).")
    add_heading(doc, "7.2 Booking Request Status", level=2)
    add_para(doc, "Pending → Approved (creates Customer + Booking) OR Rejected (unit returns to Available).")
    add_heading(doc, "7.3 Booking Status Lifecycle", level=2)
    add_para(doc, "AwaitingBookingAmount → PaymentPlanActive (when BookingAmountReceived >= BookingAmountRequired) → [PossessionGiven, SaleCompleted defined but not yet implemented in services] OR Cancelled.")
    add_heading(doc, "7.4 Installment Status", level=2)
    add_para(doc, "Pending → PartiallyPaid or Paid (on payment). Overdue computed at read time when DueDate passed and not fully paid.")

    # 8. Services
    add_heading(doc, "8. Backend Services and Key Methods")
    add_table(
        doc,
        ["Service", "Key Methods", "Responsibility"],
        [
            ["AuthService", "RegisterAsync, Login, RefreshToken", "User registration, JWT + refresh token"],
            ["TokenService", "GenerateAccessToken, GenerateRefreshToken", "JWT claims: UserId, Email, Name, Role"],
            ["BookingRequestService", "CreateBookingRequestAsync, ApproveBookingRequestAsync, RejectBookingRequestAsync, GetBookingRequestsAsync", "Online request workflow"],
            ["BookingService", "CreateBookingAsync, CreateBookingForApprovedRequestAsync, UpdateBookingFinancialsAsync, RecordBookingAmountPaymentAsync, CancelBookingAsync, GetPaymentReceiptAsync", "Confirmed booking lifecycle"],
            ["InstallmentService", "GenerateScheduleAsync, RecordInstallmentPaymentAsync, GetScheduleAsync", "Installment plan and payments"],
            ["CustomerService", "CreateCustomerAsync, FindOrCreateCustomerAsync (CNIC→Phone→Email dedupe)", "Customer CRM"],
            ["ProjectService", "CreateProjectAsync, UpdateProjectAsync, GetAllProjectsAsync", "Project CRUD"],
            ["UnitService", "CreateUnitAsync, UpdateUnitAsync, DeleteUnitAsync, GetUnitsByProjectIdAsync", "Unit CRUD"],
            ["MediaService", "Upload, Delete, Reorder, SetCover for Project/Unit media", "File management"],
            ["EmployeeService", "RecordAttendanceAsync, GenerateSalaryAsync (creates Expense), AssignTaskAsync", "HR operations"],
            ["FinanceService", "GetDashboardAsync, ManualRevenue CRUD, Expense CRUD", "Financial reporting"],
        ],
        [1.3, 2.5, 2.0],
    )

    # 9. API Endpoints
    add_heading(doc, "9. Complete API Endpoint Reference")
    endpoints = [
        ["POST", "/api/Auth/register", "Anonymous", "Register client user"],
        ["POST", "/api/Auth/login", "Anonymous", "Login → JWT + refresh token"],
        ["POST", "/api/Auth/refresh", "Anonymous", "Rotate refresh token"],
        ["GET", "/api/Auth/profile", "Authenticated", "Get userId, email, role"],
        ["POST", "/api/BookingRequest", "Anonymous/Auth", "Submit booking request"],
        ["GET", "/api/BookingRequest", "Admin", "List requests (filter, paginate)"],
        ["GET", "/api/BookingRequest/my-requests", "Auth", "Client's own requests"],
        ["POST", "/api/BookingRequest/{id}/approve", "Admin", "Approve → Customer + Booking"],
        ["POST", "/api/BookingRequest/{id}/reject", "Admin", "Reject request"],
        ["GET", "/api/BookingRequest/stats", "Admin", "Pending/approved/rejected counts"],
        ["POST", "/api/Booking", "Admin", "Create direct/walk-in booking"],
        ["GET", "/api/Booking", "Admin", "List confirmed bookings"],
        ["GET", "/api/Booking/{id}", "Admin", "Booking detail"],
        ["PUT", "/api/Booking/{id}/financials", "Admin", "Set sale price, discount, booking amount"],
        ["POST", "/api/Booking/{id}/booking-amount-payment", "Admin", "Record booking amount payment"],
        ["POST", "/api/Booking/{id}/installment-plan/generate", "Admin", "Generate/regenerate installment schedule"],
        ["POST", "/api/Booking/{id}/installments/{installmentId}/payment", "Admin", "Record installment payment"],
        ["GET", "/api/Booking/{id}/installments", "Admin", "Get installment schedule"],
        ["GET", "/api/Booking/{id}/payments", "Admin", "List all payments"],
        ["GET", "/api/Booking/{id}/payments/{paymentId}/receipt", "Admin", "Payment receipt DTO"],
        ["POST", "/api/Booking/{id}/cancel", "Admin", "Cancel booking"],
        ["GET", "/api/MyProjects", "Client", "Client's bookings by email match"],
        ["GET", "/api/MyProjects/{id}", "Client", "Client booking detail"],
        ["GET", "/api/MyProjects/{id}/installments", "Client", "Client installment schedule"],
        ["GET", "/api/MyProjects/{id}/payments/{paymentId}/receipt", "Client", "Client receipt"],
        ["GET/POST/PUT", "/api/Customer", "Admin", "Customer list/create/update"],
        ["GET/POST/PUT", "/api/Project", "Mixed", "Project CRUD (read public, write admin)"],
        ["GET/POST/PUT/DELETE", "/api/Unit", "Mixed", "Unit CRUD"],
        ["GET/POST/PUT/DELETE", "/api/Project/{id}/media", "Mixed", "Project media"],
        ["GET/POST/PUT/DELETE", "/api/Unit/{id}/media", "Mixed", "Unit media"],
        ["GET/POST/PUT/DELETE", "/api/Employee", "Admin", "Employee CRUD"],
        ["POST", "/api/Employee/{id}/attendance", "Admin", "Record attendance"],
        ["GET", "/api/Employee/attendance/batch?date=", "Admin", "Daily attendance sheet"],
        ["POST", "/api/Employee/{id}/salary", "Admin", "Generate salary + expense"],
        ["GET", "/api/Employee/salary/batch?month&year=", "Admin", "Monthly payroll batch"],
        ["GET", "/api/Finance/dashboard", "Admin", "Financial summary + revenue/expense lines"],
        ["POST/PUT/DELETE", "/api/Finance/revenue", "Admin", "Manual revenue CRUD"],
        ["POST/PUT/DELETE", "/api/Finance/expenses", "Admin", "Expense CRUD"],
    ]
    add_table(doc, ["Method", "Route", "Auth", "Purpose"], endpoints, [0.6, 2.8, 0.9, 2.5])

    # 10. Frontend routes
    add_heading(doc, "10. Frontend Routes and Pages")
    add_table(
        doc,
        ["Route", "Page", "Roles", "Primary API calls"],
        [
            ["/", "HomePage", "All", "None"],
            ["/projects", "ProjectsPage", "All", "GET /api/Project"],
            ["/projects/:id", "ProjectDetailPage", "All", "GET /api/Project/:id, GET /api/Unit/project/:id"],
            ["/units/:id", "UnitDetailPage", "All", "GET /api/Unit/:id, POST /api/BookingRequest"],
            ["/my-projects", "MyProjectsPage", "Client", "GET /api/MyProjects"],
            ["/my-projects/:id", "MyProjectDetailPage", "Client", "GET /api/MyProjects/:id, GET installments"],
            ["/bookings", "BookingRequestsPage", "Admin", "GET /api/BookingRequest, approve/reject"],
            ["/confirmed-bookings", "ConfirmedBookingsPage", "Admin", "GET /api/Booking"],
            ["/confirmed-bookings/new", "CreateBookingPage", "Admin", "POST /api/Booking"],
            ["/confirmed-bookings/:id", "BookingDetailPage", "Admin", "Booking, installments, payments APIs"],
            ["/application-form", "ApplicationFormPage", "Admin", "GET /api/Booking/:id"],
            ["/receipt/:bookingId/:paymentId", "ReceiptPage", "Client/Admin", "Receipt API (role-specific path)"],
            ["/customers", "CustomersPage", "Admin", "GET /api/Customer"],
            ["/employees", "EmployeesPage", "Admin", "Employee, attendance, salary APIs"],
            ["/finance", "FinanceDashboardPage", "Admin", "GET /api/Finance/dashboard"],
        ],
        [1.5, 1.5, 0.8, 2.0],
    )

    # 11. Recommended SSDs
    add_heading(doc, "11. Recommended System Sequence Diagrams to Create")
    add_para(
        doc,
        "Create one SSD per use case below. Use actors from Section 3. Show synchronous "
        "request/response messages between UI, Controller, Service, and Database. Include "
        "return values and state changes in notes or message labels.",
    )
    recommended = [
        "SSD-01: User Registration (Guest → AuthController → AuthService → Database)",
        "SSD-02: User Login and Profile Load (Guest → AuthModal → AuthController → TokenService → Database → GET profile)",
        "SSD-03: Token Refresh on 401 (Frontend api.ts → AuthController → AuthService)",
        "SSD-04: Browse Projects and Units (Guest → ProjectsPage → ProjectController/UnitController → Database)",
        "SSD-05: Submit Booking Request (Guest/Client → UnitDetailPage → BookingRequestController → BookingRequestService → Database; Unit→PendingReview)",
        "SSD-06: Admin Approve Booking Request (Admin → BookingRequestsPage → approve → CustomerService.FindOrCreate → BookingService.CreateBooking → Database)",
        "SSD-07: Admin Reject Booking Request (Admin → reject → Unit→Available)",
        "SSD-08: Admin Create Walk-in Booking (Admin → CreateBookingPage → BookingController → BookingService → Database)",
        "SSD-09: Set Booking Financial Terms (Admin → BookingDetailPage → PUT financials → BookingService)",
        "SSD-10: Record Booking Amount Payment (Admin → POST booking-amount-payment → Payment created → Status→PaymentPlanActive → Unit→OnPaymentPlan)",
        "SSD-11: Generate Installment Plan (Admin → POST installment-plan/generate → InstallmentService → Installment rows created)",
        "SSD-12: Record Installment Payment (Admin → POST installments/{id}/payment → Payment + Installment status update)",
        "SSD-13: View/Print Payment Receipt (Admin/Client → ReceiptPage → GetPaymentReceiptAsync)",
        "SSD-14: Client View My Projects (Client → MyProjectsPage → MyProjectsController → match by email)",
        "SSD-15: Cancel Booking (Admin → POST cancel → Booking Cancelled, Unit Available)",
        "SSD-16: Mark Employee Attendance (Admin → EmployeesAttendancePanel → EmployeeController → EmployeeService)",
        "SSD-17: Generate Employee Salary (Admin → POST salary → EmployeeSalary + Expense created)",
        "SSD-18: Finance Dashboard Load (Admin → FinanceController → FinanceService aggregates Payments + ManualRevenue + Expenses)",
        "SSD-19: Add Manual Revenue (Admin → POST /api/Finance/revenue)",
        "SSD-20: Add Expense (Admin → POST /api/Finance/expenses)",
        "SSD-21: Upload Project Media (Admin → MediaService → LocalFileStorageService → Database)",
    ]
    add_bullets(doc, recommended)

    # 12. Detailed sequences
    add_heading(doc, "12. Detailed Step-by-Step Sequences for Each Major SSD")

    add_ssd_sequence(
        doc,
        "SSD-01: User Registration",
        "Guest, Frontend (AuthModal), AuthController, AuthService, Database",
        "Guest is not logged in; email not already registered.",
        [
            "Guest clicks Sign Up in navbar.",
            "Frontend opens AuthModal with fullName, email, password fields.",
            "Guest submits form.",
            "Frontend sends POST /api/Auth/register { fullName, email, password, roleId: 2 } (no bearer token).",
            "AuthController calls AuthService.RegisterAsync.",
            "AuthService validates input; checks email uniqueness in Users table.",
            "AuthService loads Role where Role_name = Client.",
            "AuthService hashes password with BCrypt.",
            "AuthService inserts new User row into Database.",
            "Database confirms save.",
            "AuthService returns success to AuthController.",
            "AuthController returns HTTP 200 to Frontend.",
            "Frontend shows success alert and prompts user to log in.",
        ],
        "New Client user exists in Users table.",
    )

    add_ssd_sequence(
        doc,
        "SSD-02: User Login",
        "User, Frontend (AuthModal, App), AuthController, AuthService, TokenService, Database",
        "User has registered account.",
        [
            "User clicks Log in and submits email + password.",
            "Frontend POST /api/Auth/login { email, password }.",
            "AuthController → AuthService.Login.",
            "AuthService finds User by email in Database.",
            "AuthService verifies password with BCrypt.Verify.",
            "TokenService.GenerateAccessToken (JWT, 15 min, claims: UserId, Email, Name, Role).",
            "TokenService.GenerateRefreshToken (64-byte random, stored on User, 15-day expiry).",
            "AuthService updates User refresh token in Database.",
            "AuthController returns { accessToken, refreshToken, expiresInMinutes }.",
            "Frontend stores tokens in localStorage.",
            "Frontend calls GET /api/Auth/profile with Bearer token.",
            "AuthController returns { userId, email, role }.",
            "Frontend updates navbar (Admin gets Requests/Bookings/Customers/Employees/Finance links).",
        ],
        "User session active; role-based navigation visible.",
    )

    add_ssd_sequence(
        doc,
        "SSD-05: Submit Booking Request",
        "Guest/Client, UnitDetailPage, BookingRequestModal, BookingRequestController, BookingRequestService, Database",
        "Unit.Status = Available; no other Pending request on same unit.",
        [
            "User navigates to /units/:id and clicks Request Booking.",
            "Frontend opens BookingRequestModal.",
            "User enters fullName, phone, email, CNIC, address, notes.",
            "Frontend POST /api/BookingRequest { unitId, fullName, phone, email, cnic, address, notes } (no auth required).",
            "BookingRequestController → BookingRequestService.CreateBookingRequestAsync.",
            "Service validates unit exists and Status == Available.",
            "Service checks no Pending BookingRequest for unit.",
            "Service inserts BookingRequest (Status=Pending) into Database.",
            "Service updates Unit.Status = PendingReview in Database.",
            "Service returns BookingRequestResponseDto.",
            "Frontend shows success; unit card shows PendingReview.",
        ],
        "BookingRequest=Pending; Unit=PendingReview.",
    )

    add_ssd_sequence(
        doc,
        "SSD-06: Admin Approve Booking Request",
        "Admin, BookingRequestsPage, BookingRequestController, BookingRequestService, CustomerService, BookingService, Database",
        "BookingRequest.Status = Pending.",
        [
            "Admin opens /bookings and views pending request list (GET /api/BookingRequest).",
            "Admin clicks Approve on request id.",
            "Frontend POST /api/BookingRequest/{id}/approve (Admin JWT).",
            "BookingRequestService.ApproveBookingRequestAsync loads request + unit.",
            "CustomerService.FindOrCreateCustomerAsync: search by CNIC, then Phone, then Email; create if not found.",
            "Service sets BookingRequest.Status = Approved, ReviewedAt, ReviewedByUserId, CustomerId.",
            "BookingService.CreateBookingForApprovedRequestAsync: create Booking (Status=AwaitingBookingAmount, BookingReference=BK-{id:D6}).",
            "Service sets Unit.Status = Reserved.",
            "All changes saved to Database.",
            "Frontend refreshes request list and stats.",
        ],
        "Customer exists; Booking created; Unit=Reserved; Request=Approved.",
    )

    add_ssd_sequence(
        doc,
        "SSD-08: Admin Create Walk-in Booking",
        "Admin, CreateBookingPage, BookingController, BookingService, CustomerService, Database",
        "Selected unit Status = Available.",
        [
            "Admin navigates /confirmed-bookings/new.",
            "Frontend loads GET /api/Project and GET /api/Customer.",
            "Admin selects project → GET /api/Unit/project/{projectId} (Available units).",
            "Admin fills customer info, financial terms, application form fields, next-of-kin.",
            "Frontend POST /api/Booking with full payload.",
            "BookingService.CreateBookingAsync validates unit available.",
            "CustomerService resolves existing customerId or FindOrCreateCustomerAsync.",
            "BookingService creates Booking (AwaitingBookingAmount), assigns BK reference.",
            "Unit.Status set to Reserved.",
            "Database save; return BookingResponseDto.",
            "Frontend navigates to /application-form with booking data.",
        ],
        "Direct booking created without BookingRequest.",
    )

    add_ssd_sequence(
        doc,
        "SSD-10: Record Booking Amount Payment",
        "Admin, BookingDetailPage, BookingController, BookingService, Database",
        "Booking.Status = AwaitingBookingAmount; BookingAmountRequired > 0.",
        [
            "Admin opens /confirmed-bookings/:id.",
            "Frontend loads GET booking, installments, payments.",
            "Admin clicks Record Payment for booking amount.",
            "Admin enters amount, paymentMethod, reference, paidAt, notes.",
            "Frontend POST /api/Booking/{id}/booking-amount-payment.",
            "BookingService validates amount <= remaining booking amount.",
            "Service creates Payment { Type=BookingAmount, ReceiptNumber=RCP-XXXXXX }.",
            "Service increments BookingAmountReceived.",
            "If BookingAmountReceived >= BookingAmountRequired: Booking.Status = PaymentPlanActive; BookingAmountConfirmedDate = now; Unit.Status = OnPaymentPlan.",
            "Database save; return updated booking.",
            "Frontend reloads financial summary and payment history.",
        ],
        "Booking may transition to PaymentPlanActive; receipt number assigned.",
    )

    add_ssd_sequence(
        doc,
        "SSD-11: Generate Installment Plan",
        "Admin, BookingDetailPage, BookingController, InstallmentService, Database",
        "Booking.Status = PaymentPlanActive; booking amount fully received.",
        [
            "Admin fills Regenerate Installment Plan form: agreedSalePrice, discountAmount, numberOfInstallments, frequency, startDate, possessionAmount.",
            "Frontend POST /api/Booking/{id}/installment-plan/generate.",
            "InstallmentService.GenerateScheduleAsync validates inputs.",
            "Service calculates installmentPool = agreedSalePrice - bookingAmountReceived - possessionAmount.",
            "If regenerating: require regenerate=true and no payments on existing installments.",
            "Service deletes old installment rows if regenerating.",
            "Service creates Installment rows: Sequence 0 (Possession) if possessionAmount > 0; Sequences 1..N (Regular) with evenly split pool.",
            "Due dates computed from frequency (Monthly +N months, etc.).",
            "Service updates Booking plan metadata.",
            "Database save; return InstallmentScheduleDto.",
            "Frontend displays schedule table with due dates and amounts.",
        ],
        "Installment schedule exists; booking plan fields updated.",
    )

    add_ssd_sequence(
        doc,
        "SSD-12: Record Installment Payment",
        "Admin, BookingDetailPage, BookingController, InstallmentService, Database",
        "Booking.Status = PaymentPlanActive; installment not fully paid.",
        [
            "Admin clicks Record Payment on installment row.",
            "Modal pre-fills amount with remainingBalance.",
            "Frontend POST /api/Booking/{id}/installments/{installmentId}/payment { amount, paymentMethod, reference, paidAt, notes }.",
            "InstallmentService.RecordInstallmentPaymentAsync validates booking and installment.",
            "Service creates Payment { Type=Installment, InstallmentId, ReceiptNumber }.",
            "Service updates Installment.Status to Paid or PartiallyPaid.",
            "Database save.",
            "Automatic revenue appears in Finance dashboard (Payment source).",
            "Frontend reloads schedule; Admin can open receipt via /receipt/:bookingId/:paymentId.",
        ],
        "Installment partially or fully paid; receipt generated.",
    )

    add_ssd_sequence(
        doc,
        "SSD-14: Client View My Projects",
        "Client, MyProjectsPage, MyProjectsController, BookingService, Database",
        "Client logged in; Customer record exists with same email as JWT.",
        [
            "Client navigates /my-projects.",
            "Frontend GET /api/MyProjects (Client JWT).",
            "MyProjectsController extracts email from JWT claims.",
            "BookingService.GetBookingsByCustomerEmailAsync queries Bookings where Customer.Email matches.",
            "Database returns booking list with project, unit, payment summaries.",
            "Frontend renders booking cards with status and amounts.",
            "Client clicks booking → /my-projects/:id → GET /api/MyProjects/{id} and installments.",
            "Client views read-only payment history and schedule (no record payment buttons).",
        ],
        "Client sees only own bookings matched by email.",
    )

    add_ssd_sequence(
        doc,
        "SSD-17: Generate Employee Salary",
        "Admin, EmployeesSalaryPanel, EmployeeController, EmployeeService, Database",
        "Employee not Terminated; no duplicate salary for same month/year.",
        [
            "Admin opens Employees → Salaries tab; selects month/year.",
            "Frontend GET /api/Employee/salary/batch?month=&year=.",
            "Admin marks employee Paid.",
            "Frontend POST /api/Employee/{employeeId}/salary { amount, payDate, notes }.",
            "EmployeeService.GenerateSalaryAsync validates no duplicate.",
            "Service resolves optional ProjectId from employee's latest task.",
            "Service creates Expense { Category=Salary, Amount, Vendor=employee name }.",
            "Service creates EmployeeSalary linked to ExpenseId.",
            "Database save.",
            "Admin can View Receipt (salary slip); expense appears in Finance dashboard.",
        ],
        "EmployeeSalary and linked Expense created.",
    )

    add_ssd_sequence(
        doc,
        "SSD-18: Finance Dashboard",
        "Admin, FinanceDashboardPage, FinanceController, FinanceService, Database",
        "Admin authenticated.",
        [
            "Admin navigates /finance.",
            "Frontend GET /api/Finance/dashboard?projectId=&from=&to=.",
            "FinanceService.GetDashboardAsync queries Payments on non-cancelled bookings (automatic revenue).",
            "Service queries ManualRevenue records (manual revenue).",
            "Service queries Expense records (includes salary expenses).",
            "Service computes: TotalRevenue, TotalExpenses, NetProfit, OutstandingAmount (AgreedSalePrice - Payments), OverdueAmount (past-due installments).",
            "Return FinanceDashboardDto with summary cards and line items.",
            "Frontend renders Revenue/Expenses tabs; manual rows editable.",
        ],
        "Dashboard shows combined automatic + manual financial picture.",
    )

    # 13. Business rules
    add_heading(doc, "13. Important Business Rules")
    rules = [
        "Booking reference format: BK-{Id:D6} (e.g. BK-000007). Receipt format: RCP-{sequential}.",
        "Customer deduplication order: CNIC first, then Phone, then Email.",
        "Booking requests do not require authentication; optional UserId attached if logged in.",
        "Partial booking amount payments supported before PaymentPlanActive transition.",
        "Installment pool = AgreedSalePrice - BookingAmountReceived - PossessionAmount.",
        "Regenerating installment plan requires regenerate=true and no payments on existing installments.",
        "Overdue installment status is computed at read time (no background scheduler).",
        "Salary payment automatically creates Expense with Category=Salary.",
        "Finance automatic revenue comes from Payment records; manual revenue from ManualRevenue table.",
        "Client MyProjects uses email match, not Customer.UserId linkage.",
        "PossessionGiven and SaleCompleted booking statuses exist in enum but are not yet implemented in service transitions.",
    ]
    add_bullets(doc, rules)

    # 14. Entity relationship summary
    add_heading(doc, "14. Entity Relationship Summary")
    add_para(
        doc,
        "Role 1──* User. User 0..1──* Customer. Customer 1──* Booking. Project 1──* Unit. "
        "Project 1──* ProjectMedia. Unit 1──* UnitMedia. Unit 1──* Booking. Unit 1──* BookingRequest. "
        "BookingRequest 0..1── Booking. Booking 1──* Installment 1──* Payment. Booking 1──* Payment. "
        "Employee 1──* EmployeeAttendance, EmployeeTask, EmployeeSalary. EmployeeSalary *──0..1 Expense. "
        "Expense *──0..1 Project. ManualRevenue *──0..1 Project.",
    )

    # 15. Instructions for Claude
    add_heading(doc, "15. Instructions for SSD Generation Tool")
    add_bullets(
        doc,
        [
            "Draw System Sequence Diagrams (not class diagrams) — show actors, system boundary, and message flow over time.",
            "Use the recommended SSD list in Section 11 as the minimum set; Section 12 provides exact steps.",
            "Label messages with HTTP method + route for API calls (e.g. POST /api/BookingRequest).",
            "Include return messages (dashed arrows) with key response data.",
            "Show Database as a single lifeline; label operations as insert/update/query.",
            "Annotate state changes on entities (e.g. Unit.Status: Available → PendingReview).",
            "Match sample document style: one SSD per use case with figure number and caption.",
            "Primary system name: DAMS or Deen Associate Management System.",
            "Do not invent features not listed in this document.",
        ],
    )

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
