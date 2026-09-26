"""Generate DAMS test cases Word document matching sample table layout."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "docs" / "DAMS_Test_Cases.docx"

COL_WIDTHS = (1.35, 1.55, 1.35, 2.05)
GREY = "D9D9D9"


def set_cell_shading(cell, fill: str = GREY) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_pr.append(tc_w)


def set_table_fixed_layout(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def write_cell(cell, text: str, bold: bool = False, size: int = 10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)


def merge_row(table, row_idx: int, start_col: int, end_col: int) -> None:
    top = table.rows[row_idx].cells[start_col]
    for col in range(start_col + 1, end_col + 1):
        top.merge(table.rows[row_idx].cells[col])


def add_test_case(
    doc: Document,
    tc_id: str,
    name: str,
    description: str,
    actors: str,
    steps: list[tuple[str, str]],
    conditions: str,
    input_data: str,
    expected: str,
    actual: str,
    priority: str,
    frequency: str,
    acceptance: str = "Passed",
) -> None:
    table = doc.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    set_table_fixed_layout(table)
    for row in table.rows:
        for idx, width in enumerate(COL_WIDTHS):
            set_cell_width(row.cells[idx], width)

    def add_row(cells: list[tuple[str, bool, bool]]) -> None:
        row = table.add_row().cells
        for idx, width in enumerate(COL_WIDTHS):
            set_cell_width(row[idx], width)
        for i, (text, bold, grey) in enumerate(cells):
            write_cell(row[i], text, bold=bold)
            if grey:
                set_cell_shading(row[i])

    add_row(
        [
            ("Test case ID:", True, True),
            (tc_id, True, True),
            ("Test Case Name:", True, False),
            (name, False, False),
        ]
    )
    add_row(
        [
            ("Test Case Description:", True, False),
            (description, False, False),
            ("", False, False),
            ("", False, False),
        ]
    )
    merge_row(table, 1, 1, 3)
    add_row(
        [
            ("Primary Actor:", True, False),
            (actors, False, False),
            ("", False, False),
            ("", False, False),
        ]
    )
    merge_row(table, 2, 1, 3)

    header = table.add_row().cells
    for idx, width in enumerate(COL_WIDTHS):
        set_cell_width(header[idx], width)
    header[0].merge(header[3])
    write_cell(header[0], "Main success scenario:", bold=True)
    set_cell_shading(header[0])

    sub = table.add_row().cells
    sub[0].merge(sub[1])
    sub[2].merge(sub[3])
    for idx, width in enumerate(COL_WIDTHS):
        set_cell_width(sub[idx], width)
    write_cell(sub[0], "User Action", bold=True)
    set_cell_shading(sub[0])
    write_cell(sub[2], "System Response", bold=True)
    set_cell_shading(sub[2])

    step_no = 1
    for user_action, system_response in steps:
        row = table.add_row().cells
        for idx, width in enumerate(COL_WIDTHS):
            set_cell_width(row[idx], width)
        row[0].merge(row[1])
        row[2].merge(row[3])
        if user_action:
            write_cell(row[0], f"{step_no}-\n{user_action}")
            step_no += 1
        if system_response:
            write_cell(row[2], f"{step_no}-\n{system_response}")
            step_no += 1

    req = table.add_row().cells
    req[0].merge(req[3])
    write_cell(req[0], "Testing Requirement:", bold=True)
    set_cell_shading(req[0])

    add_row([("Testing Condition:", True, False), (conditions, False, False), ("", False, False), ("", False, False)])
    merge_row(table, len(table.rows) - 1, 1, 3)
    add_row([("Input Data:", True, False), (input_data, False, False), ("", False, False), ("", False, False)])
    merge_row(table, len(table.rows) - 1, 1, 3)
    add_row([("Expected Result:", True, False), (expected, False, False), ("", False, False), ("", False, False)])
    merge_row(table, len(table.rows) - 1, 1, 3)
    add_row([("Actual Result:", True, False), (actual, False, False), ("", False, False), ("", False, False)])
    merge_row(table, len(table.rows) - 1, 1, 3)
    add_row(
        [
            ("Priority:", True, False),
            (priority, False, False),
            ("Frequency:", True, False),
            (frequency, False, False),
        ]
    )
    add_row([("Test Acceptance:", True, False), (acceptance, False, False), ("", False, False), ("", False, False)])
    merge_row(table, len(table.rows) - 1, 1, 3)

    doc.add_paragraph("")


def build_test_cases() -> list[dict]:
    system = "Deen Associate Management System (DAMS)"
    common_admin = "<Admin>"
    common_client = "<Client>"
    common_both = "<Admin> <Client>"

    def tc(
        num: int,
        name: str,
        desc: str,
        actors: str,
        steps: list[tuple[str, str]],
        conditions: str,
        input_data: str,
        expected: str,
        priority: str = "High",
        frequency: str = "Frequent",
    ) -> dict:
        return {
            "id": f"TC-{num:02d}",
            "name": name,
            "description": f"This test case describes that {desc} in the {system}.",
            "actors": actors,
            "steps": steps,
            "conditions": conditions,
            "input_data": input_data,
            "expected": expected,
            "actual": expected,
            "priority": priority,
            "frequency": frequency,
        }

    return [
        tc(
            1,
            "Sign Up",
            "a new user can register an account",
            common_both,
            [
                ("User opens the sign up form and enters name, email, phone, password, and role.", ""),
                ("", "System validates the fields and creates the account successfully."),
                ("Steps repeat if another user registers.", ""),
            ],
            "System must be in running state.\nEmail and phone must not already exist.",
            "Name, email, phone, password, role",
            "User account is created successfully.",
        ),
        tc(
            2,
            "Login",
            "a registered user can log in",
            common_both,
            [
                ("User enters email and password on the login page.", ""),
                ("", "System verifies credentials and redirects the user to the dashboard."),
            ],
            "System must be in running state.\nUser must have valid login credentials.",
            "Email and password",
            "User is successfully logged into the system.",
        ),
        tc(
            3,
            "Logout",
            "a logged-in user can log out",
            common_both,
            [
                ("User clicks the Logout button.", ""),
                ("", "System ends the session and redirects the user to the home page."),
            ],
            "System must be in running state.\nUser must be logged in.",
            "Click Logout",
            "User is logged out successfully.",
        ),
        tc(
            4,
            "Change Password",
            "a logged-in user can change password",
            common_both,
            [
                ("User enters current password and new password.", ""),
                ("", "System validates the current password and updates the new password."),
            ],
            "System must be in running state.\nUser must be logged in.",
            "Current password and new password",
            "Password is changed successfully.",
        ),
        tc(
            5,
            "Forget Password",
            "a user can reset forgotten password",
            common_both,
            [
                ("User enters registered email on forget password page.", ""),
                ("", "System sends reset instructions and allows password reset."),
            ],
            "System must be in running state.\nEmail must exist in the database.",
            "Registered email address",
            "Password reset process completes successfully.",
        ),
        tc(
            6,
            "Add Project",
            "an admin can add a new project",
            common_admin,
            [
                ("Admin enters project name, location, description, and dates.", ""),
                ("", "System validates the data and saves the new project."),
            ],
            "System must be in running state.\nAdmin must be logged in.",
            "Project name, location, description, dates",
            "New project is added successfully.",
        ),
        tc(
            7,
            "Update Project",
            "an admin can update project details",
            common_admin,
            [
                ("Admin selects a project and updates its details.", ""),
                ("", "System validates and saves the updated project information."),
            ],
            "System must be in running state.\nProject must already exist.",
            "Updated project details",
            "Project details are updated successfully.",
        ),
        tc(
            8,
            "View Project",
            "a user can view project details",
            common_both,
            [
                ("User selects a project from the projects list.", ""),
                ("", "System displays complete project details and related units."),
            ],
            "System must be in running state.",
            "Project selection",
            "Project details are displayed correctly.",
            "Medium",
        ),
        tc(
            9,
            "Search Project",
            "an admin can search projects",
            common_admin,
            [
                ("Admin enters search criteria on the projects page.", ""),
                ("", "System displays matching projects in a list."),
            ],
            "System must be in running state.\nAdmin must be logged in.",
            "Project name or location",
            "Matching projects are displayed.",
            "Medium",
            "Less frequent",
        ),
        tc(
            10,
            "Upload Project Media",
            "an admin can upload project media",
            common_admin,
            [
                ("Admin selects image or video files for a project.", ""),
                ("", "System uploads and displays the media in the project gallery."),
            ],
            "System must be in running state.\nValid media format must be selected.",
            "Project media file",
            "Project media is uploaded successfully.",
            "Medium",
        ),
        tc(
            11,
            "Add Unit",
            "an admin can add a unit to a project",
            common_admin,
            [
                ("Admin enters unit number, type, floor, size, price, and status.", ""),
                ("", "System validates and saves the new unit."),
            ],
            "System must be in running state.\nProject must exist.",
            "Unit details",
            "Unit is added successfully.",
        ),
        tc(
            12,
            "Update Unit",
            "an admin can update unit details",
            common_admin,
            [
                ("Admin selects a unit and updates its information.", ""),
                ("", "System validates and saves the updated unit details."),
            ],
            "System must be in running state.\nUnit must already exist.",
            "Updated unit details",
            "Unit details are updated successfully.",
        ),
        tc(
            13,
            "View Unit Status",
            "a user can view unit availability status",
            common_both,
            [
                ("User opens a unit detail page.", ""),
                ("", "System displays unit status such as Available, Booked, or Sold."),
            ],
            "System must be in running state.",
            "Unit selection",
            "Unit status is displayed correctly.",
            "Medium",
        ),
        tc(
            14,
            "Upload Unit Media",
            "an admin can upload unit media",
            common_admin,
            [
                ("Admin uploads images or videos for a selected unit.", ""),
                ("", "System stores and displays the unit media."),
            ],
            "System must be in running state.\nValid media format must be selected.",
            "Unit media file",
            "Unit media is uploaded successfully.",
            "Medium",
        ),
        tc(
            15,
            "Add Customer",
            "an admin can add a customer",
            common_admin,
            [
                ("Admin enters customer name, phone, CNIC, email, and address.", ""),
                ("", "System validates and saves the customer record."),
            ],
            "System must be in running state.\nPhone or CNIC must be unique.",
            "Customer details",
            "Customer is added successfully.",
        ),
        tc(
            16,
            "Update Customer",
            "an admin can update customer information",
            common_admin,
            [
                ("Admin selects a customer and updates the details.", ""),
                ("", "System validates and saves the updated customer record."),
            ],
            "System must be in running state.\nCustomer must exist.",
            "Updated customer details",
            "Customer information is updated successfully.",
        ),
        tc(
            17,
            "View Customer",
            "an admin can view customer details",
            common_admin,
            [
                ("Admin selects a customer from the customer list.", ""),
                ("", "System displays customer profile and booking history."),
            ],
            "System must be in running state.",
            "Customer selection",
            "Customer details are displayed correctly.",
            "Medium",
        ),
        tc(
            18,
            "Search Customer",
            "an admin can search customers",
            common_admin,
            [
                ("Admin enters name, phone, or CNIC in search field.", ""),
                ("", "System displays matching customer records."),
            ],
            "System must be in running state.",
            "Search keyword",
            "Matching customers are displayed.",
            "Medium",
            "Less frequent",
        ),
        tc(
            19,
            "Submit Booking Request",
            "a client can submit a booking request",
            common_client,
            [
                ("Client selects a unit and fills the booking request form.", ""),
                ("", "System saves the request with Pending status."),
            ],
            "System must be in running state.\nClient must be logged in.\nUnit must be available.",
            "Client details and unit selection",
            "Booking request is submitted successfully.",
        ),
        tc(
            20,
            "View Booking Requests",
            "an admin can view booking requests",
            common_admin,
            [
                ("Admin opens the booking requests page.", ""),
                ("", "System displays all pending and reviewed booking requests."),
            ],
            "System must be in running state.\nAdmin must be logged in.",
            "Open booking requests page",
            "Booking requests are displayed correctly.",
            "Medium",
        ),
        tc(
            21,
            "Approve Booking Request",
            "an admin can approve a booking request",
            common_admin,
            [
                ("Admin reviews a pending booking request and clicks Approve.", ""),
                ("", "System updates request status and allows booking creation."),
            ],
            "System must be in running state.\nBooking request must be pending.",
            "Booking request ID",
            "Booking request is approved successfully.",
        ),
        tc(
            22,
            "Reject Booking Request",
            "an admin can reject a booking request",
            common_admin,
            [
                ("Admin reviews a pending booking request and clicks Reject.", ""),
                ("", "System updates request status to rejected with reason."),
            ],
            "System must be in running state.\nBooking request must be pending.",
            "Booking request ID and rejection reason",
            "Booking request is rejected successfully.",
        ),
        tc(
            23,
            "Create Confirmed Booking",
            "an admin can create a confirmed booking",
            common_admin,
            [
                ("Admin selects customer, unit, and enters booking financial details.", ""),
                ("", "System creates the booking and generates booking reference."),
            ],
            "System must be in running state.\nCustomer and unit must exist.",
            "Customer, unit, and booking details",
            "Confirmed booking is created successfully.",
        ),
        tc(
            24,
            "Update Booking",
            "an admin can update booking details",
            common_admin,
            [
                ("Admin opens a booking and updates financial or status details.", ""),
                ("", "System validates and saves the updated booking."),
            ],
            "System must be in running state.\nBooking must exist.",
            "Updated booking details",
            "Booking is updated successfully.",
        ),
        tc(
            25,
            "View Booking Details",
            "an admin can view booking details",
            common_admin,
            [
                ("Admin selects a booking from the bookings list.", ""),
                ("", "System displays booking, customer, unit, and payment details."),
            ],
            "System must be in running state.",
            "Booking selection",
            "Booking details are displayed correctly.",
            "Medium",
        ),
        tc(
            26,
            "Cancel Booking",
            "an admin can cancel a booking",
            common_admin,
            [
                ("Admin selects a booking and chooses Cancel Booking.", ""),
                ("", "System updates booking status to cancelled."),
            ],
            "System must be in running state.\nBooking must be active.",
            "Booking ID and cancellation reason",
            "Booking is cancelled successfully.",
        ),
        tc(
            27,
            "View Booking History",
            "a user can view booking history",
            common_both,
            [
                ("User opens booking history or my projects page.", ""),
                ("", "System displays past and current booking records."),
            ],
            "System must be in running state.\nUser must be logged in.",
            "Open booking history",
            "Booking history is displayed correctly.",
            "Medium",
        ),
        tc(
            28,
            "Record Booking Amount Payment",
            "an admin can record booking amount payment",
            common_admin,
            [
                ("Admin opens a booking and records booking amount payment.", ""),
                ("", "System saves payment and updates booking amount received."),
            ],
            "System must be in running state.\nBooking must exist.",
            "Payment amount and method",
            "Booking amount payment is recorded successfully.",
        ),
        tc(
            29,
            "Generate Installment Plan",
            "an admin can generate installment plan",
            common_admin,
            [
                ("Admin enters installment frequency, count, and start date.", ""),
                ("", "System generates installment schedule for the booking."),
            ],
            "System must be in running state.\nBooking must be confirmed.",
            "Installment plan details",
            "Installment plan is generated successfully.",
        ),
        tc(
            30,
            "Record Installment Payment",
            "an admin can record installment payment",
            common_admin,
            [
                ("Admin selects an installment and records payment.", ""),
                ("", "System updates installment status and booking payment totals."),
            ],
            "System must be in running state.\nInstallment must be pending.",
            "Installment ID and payment amount",
            "Installment payment is recorded successfully.",
        ),
        tc(
            31,
            "View Payment History",
            "an admin can view payment history",
            common_admin,
            [
                ("Admin opens payment history for a booking.", ""),
                ("", "System displays all booking and installment payments."),
            ],
            "System must be in running state.",
            "Booking selection",
            "Payment history is displayed correctly.",
            "Medium",
        ),
        tc(
            32,
            "Generate Payment Receipt",
            "an admin can generate payment receipt",
            common_admin,
            [
                ("Admin opens a recorded payment and views receipt.", ""),
                ("", "System displays printable receipt with receipt number."),
            ],
            "System must be in running state.\nPayment must exist.",
            "Payment ID",
            "Payment receipt is generated successfully.",
            "Medium",
        ),
        tc(
            33,
            "Print Application Form",
            "an admin can print application form",
            common_admin,
            [
                ("Admin opens application form for a booking.", ""),
                ("", "System displays printable official application form."),
            ],
            "System must be in running state.\nBooking must exist.",
            "Booking ID",
            "Application form is displayed correctly.",
            "Medium",
            "Less frequent",
        ),
        tc(
            34,
            "Add Employee",
            "an admin can add an employee",
            common_admin,
            [
                ("Admin enters employee name, job title, department, phone, and salary.", ""),
                ("", "System validates and saves the employee record."),
            ],
            "System must be in running state.",
            "Employee details",
            "Employee is added successfully.",
        ),
        tc(
            35,
            "Update Employee",
            "an admin can update employee details",
            common_admin,
            [
                ("Admin selects an employee and updates the details.", ""),
                ("", "System validates and saves updated employee information."),
            ],
            "System must be in running state.\nEmployee must exist.",
            "Updated employee details",
            "Employee details are updated successfully.",
        ),
        tc(
            36,
            "View Employee",
            "an admin can view employee details",
            common_admin,
            [
                ("Admin selects an employee from the employee list.", ""),
                ("", "System displays employee profile and related records."),
            ],
            "System must be in running state.",
            "Employee selection",
            "Employee details are displayed correctly.",
            "Medium",
        ),
        tc(
            37,
            "Assign Staff to Project",
            "an admin can assign staff to a project",
            common_admin,
            [
                ("Admin assigns an employee to a selected project.", ""),
                ("", "System saves staff-project assignment successfully."),
            ],
            "System must be in running state.\nEmployee and project must exist.",
            "Employee ID and project ID",
            "Staff is assigned to project successfully.",
        ),
        tc(
            38,
            "Record Attendance",
            "an admin can record employee attendance",
            common_admin,
            [
                ("Admin records attendance status for an employee.", ""),
                ("", "System saves attendance record for the selected date."),
            ],
            "System must be in running state.\nEmployee must exist.",
            "Employee ID, date, attendance status",
            "Attendance is recorded successfully.",
            "Medium",
        ),
        tc(
            39,
            "Assign Employee Task",
            "an admin can assign a task to an employee",
            common_admin,
            [
                ("Admin creates a task with title, description, and due date.", ""),
                ("", "System assigns the task to the selected employee."),
            ],
            "System must be in running state.\nEmployee must exist.",
            "Task details",
            "Employee task is assigned successfully.",
            "Medium",
        ),
        tc(
            40,
            "Update Task Status",
            "an admin can update employee task status",
            common_admin,
            [
                ("Admin opens a task and updates its status.", ""),
                ("", "System saves the updated task status."),
            ],
            "System must be in running state.\nTask must exist.",
            "Task ID and new status",
            "Task status is updated successfully.",
            "Medium",
        ),
        tc(
            41,
            "Generate Salary",
            "an admin can generate employee salary",
            common_admin,
            [
                ("Admin selects employee and salary month/year.", ""),
                ("", "System generates salary record for the employee."),
            ],
            "System must be in running state.\nEmployee must exist.",
            "Employee ID, month, year, amount",
            "Salary is generated successfully.",
        ),
        tc(
            42,
            "View Salary History",
            "an admin can view salary history",
            common_admin,
            [
                ("Admin opens salary history for an employee.", ""),
                ("", "System displays past salary records."),
            ],
            "System must be in running state.",
            "Employee selection",
            "Salary history is displayed correctly.",
            "Medium",
            "Less frequent",
        ),
        tc(
            43,
            "Add Expense",
            "an admin can add an expense",
            common_admin,
            [
                ("Admin enters expense category, amount, vendor, and date.", ""),
                ("", "System validates and saves the expense record."),
            ],
            "System must be in running state.",
            "Expense details",
            "Expense is added successfully.",
        ),
        tc(
            44,
            "Update Expense",
            "an admin can update an expense",
            common_admin,
            [
                ("Admin selects an expense and updates its details.", ""),
                ("", "System validates and saves the updated expense."),
            ],
            "System must be in running state.\nExpense must exist.",
            "Updated expense details",
            "Expense is updated successfully.",
            "Medium",
        ),
        tc(
            45,
            "View Expense",
            "an admin can view expense records",
            common_admin,
            [
                ("Admin opens the finance or expense list page.", ""),
                ("", "System displays all expense records."),
            ],
            "System must be in running state.",
            "Open expense list",
            "Expense records are displayed correctly.",
            "Medium",
        ),
        tc(
            46,
            "Add Manual Income",
            "an admin can add manual income",
            common_admin,
            [
                ("Admin enters income type, amount, project, and date.", ""),
                ("", "System validates and saves the manual income record."),
            ],
            "System must be in running state.",
            "Manual income details",
            "Manual income is added successfully.",
        ),
        tc(
            47,
            "View Finance Dashboard",
            "an admin can view finance dashboard",
            common_admin,
            [
                ("Admin opens the finance dashboard.", ""),
                ("", "System displays income, expense, and summary reports."),
            ],
            "System must be in running state.\nAdmin must be logged in.",
            "Open finance dashboard",
            "Finance dashboard is displayed correctly.",
            "Medium",
        ),
        tc(
            48,
            "Browse Projects",
            "a client can browse available projects",
            common_client,
            [
                ("Client opens the projects page.", ""),
                ("", "System displays all available projects."),
            ],
            "System must be in running state.",
            "Open projects page",
            "Projects are displayed correctly.",
            "Medium",
            "Frequent",
        ),
        tc(
            49,
            "View Project Details",
            "a client can view project details",
            common_client,
            [
                ("Client selects a project from the list.", ""),
                ("", "System displays project information and available units."),
            ],
            "System must be in running state.",
            "Project selection",
            "Project details are displayed correctly.",
            "Medium",
        ),
        tc(
            50,
            "View My Projects",
            "a client can view my projects page",
            common_client,
            [
                ("Client opens the My Projects page.", ""),
                ("", "System displays projects and units linked to the client."),
            ],
            "System must be in running state.\nClient must be logged in.",
            "Open My Projects page",
            "Client projects are displayed correctly.",
            "Medium",
        ),
        tc(
            51,
            "Manage User Role",
            "a super admin can manage user roles",
            "<Super Admin>",
            [
                ("Super Admin updates role for a system user.", ""),
                ("", "System saves the updated role and access permissions."),
            ],
            "System must be in running state.\nSuper Admin must be logged in.",
            "User ID and role",
            "User role is updated successfully.",
            "High",
            "Less frequent",
        ),
        tc(
            52,
            "Invalid Login Attempt",
            "the system handles invalid login attempts",
            common_both,
            [
                ("User enters incorrect email or password.", ""),
                ("", "System displays an invalid credentials error message."),
            ],
            "System must be in running state.",
            "Invalid email or password",
            "Login is rejected and error message is shown.",
            "High",
            "Frequent",
        ),
        tc(
            53,
            "Duplicate Email Registration",
            "the system prevents duplicate email registration",
            common_both,
            [
                ("User tries to register with an email that already exists.", ""),
                ("", "System displays duplicate email error and blocks registration."),
            ],
            "System must be in running state.\nEmail must already exist.",
            "Existing email address",
            "Duplicate registration is prevented.",
            "High",
            "Less frequent",
        ),
        tc(
            54,
            "Duplicate Customer Entry",
            "the system prevents duplicate customer records",
            common_admin,
            [
                ("Admin tries to add a customer with existing phone or CNIC.", ""),
                ("", "System displays duplicate entry error and blocks save."),
            ],
            "System must be in running state.\nCustomer phone or CNIC must already exist.",
            "Duplicate phone or CNIC",
            "Duplicate customer entry is prevented.",
            "High",
            "Less frequent",
        ),
    ]


def build_document() -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    for case in build_test_cases():
        add_test_case(
            doc,
            tc_id=case["id"],
            name=case["name"],
            description=case["description"],
            actors=case["actors"],
            steps=case["steps"],
            conditions=case["conditions"],
            input_data=case["input_data"],
            expected=case["expected"],
            actual=case["actual"],
            priority=case["priority"],
            frequency=case["frequency"],
        )

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(f"Created {OUT_DOCX}")


if __name__ == "__main__":
    build_document()
